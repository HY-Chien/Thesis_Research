#!/usr/bin/env python3
"""Build two editable tl_sum thesis-defense slides and matching previews."""

from __future__ import annotations

import os
import sys
from pathlib import Path

PACKAGE_DIR = Path("/home/S113065528/.codex_tmp/python_pptx")
sys.path.insert(0, str(PACKAGE_DIR))

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
PPTX_OUT = OUT_DIR / "tl_sum_2slides.pptx"
PDF_OUT = OUT_DIR / "tl_sum_2slides.pdf"
PNG_OUT = [OUT_DIR / "tl_sum_slide_48.png", OUT_DIR / "tl_sum_slide_49.png"]
APPENDIX_PPTX_OUT = OUT_DIR / "tl_sum_appendix_2slides.pptx"
APPENDIX_PDF_OUT = OUT_DIR / "tl_sum_appendix_2slides.pdf"
FULL_PPTX_OUT = OUT_DIR / "tl_sum_2slides_with_appendix.pptx"
FULL_PDF_OUT = OUT_DIR / "tl_sum_2slides_with_appendix.pdf"
APPENDIX_PNG_OUT = [
    OUT_DIR / "tl_sum_appendix_top5.png",
    OUT_DIR / "tl_sum_appendix_entropy_derivation.png",
]

SLIDE_W_IN = 13.333333
SLIDE_H_IN = 7.5
PREVIEW_DPI = 144
PREVIEW_W = int(round(SLIDE_W_IN * PREVIEW_DPI))
PREVIEW_H = int(round(SLIDE_H_IN * PREVIEW_DPI))


COLORS = {
    "background": "FAFAF9",
    "white": "FFFFFF",
    "title": "762326",
    "ink": "18191A",
    "muted": "626B70",
    "line": "344E5C",
    "panel_line": "CED4D8",
    "panel_header": "EEF2F4",
    "ice": "DCE5EA",
    "ice_deep": "BFCED6",
    "blue": "3E6E8B",
    "blue_light": "AFC7D7",
    "green": "1E9C72",
    "green_dark": "137455",
    "green_light": "DCEFE7",
    "highlight": "D4E8C9",
    "orange": "C65B42",
    "orange_light": "F2D8CF",
    "gray_bar": "AEB8BE",
    "track": "E9EDEF",
}

FONT_SERIF = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
FONT_SERIF_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"
FONT_SERIF_ITALIC = "/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf"
FONT_SANS = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
FONT_SANS_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"
FONT_MATH = "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf"
FONT_MATH_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"


def rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def pil_color(hex_color: str) -> str:
    return f"#{hex_color.lstrip('#')}"


def pil_font(size_pt: float, family: str, bold: bool = False, italic: bool = False):
    px_size = max(8, int(round(size_pt * PREVIEW_DPI / 72)))
    if family == "sans":
        path = FONT_SANS_BOLD if bold else FONT_SANS
    elif family == "math":
        path = FONT_MATH_BOLD if bold else FONT_MATH
    elif italic:
        path = FONT_SERIF_ITALIC
    else:
        path = FONT_SERIF_BOLD if bold else FONT_SERIF
    return ImageFont.truetype(path, px_size)


def ppt_font_name(family: str) -> str:
    if family == "sans":
        return "Arial"
    if family == "math":
        return "Cambria Math"
    return "Times New Roman"


class SlideCanvas:
    def __init__(self, prs: Presentation):
        self.slide = prs.slides.add_slide(prs.slide_layouts[6])
        self.slide.background.fill.solid()
        self.slide.background.fill.fore_color.rgb = rgb(COLORS["background"])
        self.image = Image.new(
            "RGB", (PREVIEW_W, PREVIEW_H), pil_color(COLORS["background"])
        )
        self.draw = ImageDraw.Draw(self.image)

    @staticmethod
    def px(value_in: float) -> int:
        return int(round(value_in * PREVIEW_DPI))

    def rect(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: str,
        stroke: str | None = None,
        stroke_pt: float = 0.8,
        radius: float = 0.0,
    ):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(
            shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if stroke:
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = Pt(stroke_pt)
        else:
            shape.line.fill.background()
        box = [self.px(x), self.px(y), self.px(x + w), self.px(y + h)]
        if radius:
            self.draw.rounded_rectangle(
                box,
                radius=max(2, self.px(radius)),
                fill=pil_color(fill),
                outline=pil_color(stroke) if stroke else None,
                width=max(1, int(round(stroke_pt * PREVIEW_DPI / 72))) if stroke else 1,
            )
        else:
            self.draw.rectangle(
                box,
                fill=pil_color(fill),
                outline=pil_color(stroke) if stroke else None,
                width=max(1, int(round(stroke_pt * PREVIEW_DPI / 72))) if stroke else 1,
            )
        return shape

    def ellipse(
        self,
        x: float,
        y: float,
        w: float,
        h: float,
        fill: str,
        stroke: str | None = None,
    ):
        shape = self.slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if stroke:
            shape.line.color.rgb = rgb(stroke)
        else:
            shape.line.fill.background()
        self.draw.ellipse(
            [self.px(x), self.px(y), self.px(x + w), self.px(y + h)],
            fill=pil_color(fill),
            outline=pil_color(stroke) if stroke else None,
        )
        return shape

    def line(self, x: float, y: float, w: float, color: str, height: float = 0.012):
        self.rect(x, y, w, height, color)

    def _wrap_lines(self, text: str, font, max_width: int) -> list[str]:
        out: list[str] = []
        for explicit_line in text.split("\n"):
            if not explicit_line:
                out.append("")
                continue
            words = explicit_line.split(" ")
            current = words[0]
            for word in words[1:]:
                trial = current + " " + word
                width = self.draw.textbbox((0, 0), trial, font=font)[2]
                if width <= max_width:
                    current = trial
                else:
                    out.append(current)
                    current = word
            out.append(current)
        return out

    def text(
        self,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        size: float,
        color: str = COLORS["ink"],
        bold: bool = False,
        italic: bool = False,
        family: str = "serif",
        align: str = "left",
        valign: str = "middle",
        margin: float = 0.03,
        line_spacing: float = 1.08,
    ):
        textbox = self.slide.shapes.add_textbox(
            Inches(x), Inches(y), Inches(w), Inches(h)
        )
        tf = textbox.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(margin)
        tf.margin_right = Inches(margin)
        tf.margin_top = Inches(0.0)
        tf.margin_bottom = Inches(0.0)
        tf.vertical_anchor = {
            "top": MSO_ANCHOR.TOP,
            "middle": MSO_ANCHOR.MIDDLE,
            "bottom": MSO_ANCHOR.BOTTOM,
        }[valign]
        p = tf.paragraphs[0]
        p.alignment = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }[align]
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = text
        run.font.name = ppt_font_name(family)
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)

        font = pil_font(size, family, bold, italic)
        x0 = self.px(x + margin)
        y0 = self.px(y)
        max_width = max(1, self.px(w - 2 * margin))
        lines = self._wrap_lines(text, font, max_width)
        probe = self.draw.textbbox((0, 0), "Ag", font=font)
        line_h = max(1, int((probe[3] - probe[1]) * line_spacing))
        total_h = line_h * len(lines)
        if valign == "middle":
            cursor_y = y0 + max(0, (self.px(h) - total_h) // 2)
        elif valign == "bottom":
            cursor_y = y0 + max(0, self.px(h) - total_h)
        else:
            cursor_y = y0
        for line_text in lines:
            bbox = self.draw.textbbox((0, 0), line_text, font=font)
            line_w = bbox[2] - bbox[0]
            if align == "center":
                cursor_x = x0 + max(0, (max_width - line_w) // 2)
            elif align == "right":
                cursor_x = x0 + max(0, max_width - line_w)
            else:
                cursor_x = x0
            self.draw.text(
                (cursor_x, cursor_y - bbox[1]),
                line_text,
                font=font,
                fill=pil_color(color),
            )
            cursor_y += line_h
        return textbox

    def header(self, title: str, section: str, slide_no: str, title_size: float = 28.5):
        self.ellipse(12.62, -0.38, 1.42, 1.42, COLORS["ice_deep"])
        self.rect(10.82, 0.0, 2.51, 0.58, COLORS["panel_header"])
        self.text(
            section,
            10.92,
            0.02,
            2.12,
            0.48,
            18.0,
            bold=True,
            align="center",
        )
        self.text(
            title,
            0.58,
            0.26,
            10.15,
            0.62,
            title_size,
            color=COLORS["title"],
            bold=True,
        )
        self.line(0.58, 1.04, 12.05, COLORS["line"], 0.011)
        self.line(0.62, 7.12, 12.02, COLORS["line"], 0.011)
        self.ellipse(12.88, 6.43, 1.02, 1.02, "4F8191")
        self.text(
            slide_no,
            12.70,
            7.14,
            0.43,
            0.22,
            11.5,
            color=COLORS["line"],
            align="right",
        )

    def step_header(self, number: str, label: str, x: float, y: float, w: float):
        self.ellipse(x, y, 0.36, 0.36, COLORS["title"])
        self.text(number, x, y - 0.005, 0.36, 0.36, 14, color="FFFFFF", bold=True, align="center")
        self.text(label, x + 0.48, y - 0.035, w - 0.48, 0.43, 18, bold=True)


def add_slide_1(prs: Presentation) -> SlideCanvas:
    c = SlideCanvas(prs)
    c.header("Discovering the Signal: It Lives in the Tail", "Analysis", "48", 28.5)
    c.text(
        "42-signal survey  →  entropy decomposition",
        0.66,
        1.12,
        12.0,
        0.34,
        15.5,
        color=COLORS["muted"],
        italic=True,
    )

    c.rect(0.55, 1.55, 5.86, 4.72, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)
    c.rect(6.62, 1.55, 6.16, 4.72, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)

    c.step_header("1", "Broad survey", 0.84, 1.77, 4.95)
    c.text(
        "42 confidence signals  ·  4 datasets × 1,000 questions",
        0.86,
        2.22,
        5.20,
        0.33,
        13.2,
        family="sans",
        color=COLORS["muted"],
    )
    c.rect(0.86, 2.57, 5.18, 0.67, COLORS["panel_header"], radius=0.06)
    c.text(
        "H_α(p) = 1/(1−α) · log ∑ᵥ p(v)ᵅ",
        0.98,
        2.59,
        4.94,
        0.31,
        14.0,
        family="math",
        bold=True,
        align="center",
    )
    c.text(
        "α > 1: head-heavy   ·   α → 1: Shannon   ·   α < 1: tail-sensitive",
        0.97,
        2.91,
        4.96,
        0.21,
        9.7,
        family="sans",
        color=COLORS["muted"],
        align="center",
    )
    c.text("2Wiki AUROC ↑", 0.98, 3.30, 1.48, 0.25, 10.7, color=COLORS["muted"], family="sans")

    baseline_y = 5.27
    c.line(0.98, baseline_y, 4.96, COLORS["panel_line"], 0.012)
    bars = [
        (1.22, 0.72, 0.651, "BAC4CA", "α = 2", "p² → head"),
        (2.98, 1.12, 0.671, "7892A1", "α → 1", "Shannon limit"),
        (4.74, 1.72, 0.716, COLORS["green"], "α = 0.5", "√p → tail"),
    ]
    for x, height, value, fill, top_label, bottom_label in bars:
        c.rect(x, baseline_y - height, 0.72, height, fill, radius=0.04)
        c.text(f"{value:.3f}", x - 0.12, baseline_y - height - 0.34, 0.96, 0.28, 14.5, bold=True, align="center", family="sans")
        c.text(top_label, x - 0.22, 5.34, 1.16, 0.27, 12.3, bold=True, align="center", family="sans")
        c.text(bottom_label, x - 0.34, 5.61, 1.40, 0.23, 10.6, color=COLORS["muted"], align="center", family="sans")
    c.rect(1.03, 5.93, 4.88, 0.24, COLORS["green_light"], radius=0.08)
    c.text(
        "Lower α gives rare tokens more influence  →  better separation",
        1.04,
        5.91,
        4.86,
        0.27,
        10.8,
        color=COLORS["green_dark"],
        bold=True,
        family="sans",
        align="center",
    )

    c.step_header("2", "Decompose entropy", 6.91, 1.77, 5.20)
    c.rect(6.95, 2.18, 5.50, 1.34, COLORS["panel_header"], radius=0.07)
    c.text("A = top-5,  B = tail,  m₅ = P(A)", 7.10, 2.21, 5.20, 0.20, 9.8, color=COLORS["muted"], family="math", align="center")
    c.text("p_top(v)=p(v)/m₅ on A;  p_tail(v)=p(v)/(1−m₅) on B", 7.06, 2.42, 5.28, 0.22, 10.0, family="math", align="center")
    c.text("H(p) = −∑ᵥ∈A p(v)log p(v) − ∑ᵥ∈B p(v)log p(v)", 7.07, 2.66, 5.26, 0.24, 11.2, family="math", align="center")
    c.text("Substitute p=m₅p_top on A and p=(1−m₅)p_tail on B; expand log", 7.04, 2.91, 5.32, 0.18, 8.9, color=COLORS["muted"], family="sans", align="center")
    c.rect(7.12, 3.14, 5.16, 0.29, COLORS["green_light"], radius=0.05)
    c.text("⇒ H = H_B(m₅) + m₅H(p_top) + (1−m₅)H(p_tail)", 7.15, 3.12, 5.10, 0.32, 12.3, color=COLORS["green_dark"], family="math", bold=True, align="center")

    c.text("Diagnostic probes isolate shape: H(p_tail) normalizes away tail amount", 7.01, 3.57, 5.38, 0.20, 9.4, color=COLORS["muted"], family="sans", align="center")
    track_x, track_w = 9.33, 2.42
    c.text("chance = 0.50", track_x, 3.78, 1.45, 0.20, 9.7, color=COLORS["muted"], family="sans")
    c.text("macro AUROC ↑", 11.15, 3.78, 1.18, 0.20, 9.7, color=COLORS["muted"], family="sans", align="right")
    components = [
        ("Tail mass: how much escapes", "H_B(m₅)", 0.742, COLORS["green"], 4.06),
        ("Head shape: how top-5 splits", "H(p_top) · conditional", 0.712, COLORS["blue"], 4.76),
        ("Tail shape: how tail splits", "H(p_tail) · conditional", 0.502, COLORS["orange"], 5.46),
    ]
    for label, formula, value, fill, y in components:
        c.text(label, 6.96, y - 0.10, 2.24, 0.27, 11.3, bold=(value == 0.742), family="sans")
        c.text(formula, 6.96, y + 0.17, 2.24, 0.22, 9.8, color=COLORS["muted"], family="math")
        c.rect(track_x, y, track_w, 0.29, COLORS["track"], radius=0.05)
        bar_w = max(0.045, (value - 0.50) / 0.25 * track_w)
        c.rect(track_x, y, bar_w, 0.29, fill, radius=0.05)
        c.text(f"{value:.3f}", 11.84, y - 0.05, 0.58, 0.34, 12.8, bold=(value == 0.742), family="sans", align="right")
    c.rect(7.11, 5.92, 5.18, 0.26, COLORS["green_light"], radius=0.06)
    c.text(
        "Amount matters (.742); allocation inside the tail does not (.502)",
        7.14,
        5.90,
        5.12,
        0.29,
        10.5,
        color=COLORS["green_dark"],
        bold=True,
        family="sans",
        align="center",
    )

    c.rect(0.72, 6.44, 11.82, 0.50, COLORS["highlight"], radius=0.07)
    c.text(
        "Takeaway: tail-sensitive signals detect how much mass escapes the head—not how that tail mass is arranged.",
        0.92,
        6.45,
        11.40,
        0.46,
        15.0,
        color=COLORS["green_dark"],
        bold=True,
        align="center",
    )
    return c


def add_token_card(c: SlideCanvas, x: float, label: str, suffix: str, top_mass: float):
    c.rect(x, 3.73, 1.36, 1.44, COLORS["panel_header"], COLORS["panel_line"], 0.7, 0.09)
    c.text(label, x + 0.08, 3.82, 1.20, 0.27, 12.4, bold=True, family="sans", align="center")
    bar_x, bar_y, bar_w = x + 0.17, 4.24, 1.02
    c.rect(bar_x, bar_y, bar_w, 0.24, COLORS["orange_light"], radius=0.04)
    c.rect(bar_x, bar_y, bar_w * top_mass, 0.24, COLORS["green"], radius=0.04)
    c.text("top-5", bar_x, 4.51, 0.57, 0.19, 9.6, color=COLORS["green_dark"], family="sans")
    c.text("tail", bar_x + 0.60, 4.51, 0.42, 0.19, 9.6, color=COLORS["orange"], family="sans", align="right")
    c.text(f"−log m₅{suffix}", x + 0.05, 4.76, 1.26, 0.30, 12.6, family="math", align="center")


def add_slide_2(prs: Presentation) -> SlideCanvas:
    c = SlideCanvas(prs)
    c.header("From Entropy to tl_sum", "Analysis", "49", 29.0)
    c.text(
        "isolate the escape signal  →  make it monotonic  →  aggregate across tokens",
        0.66,
        1.12,
        12.0,
        0.34,
        15.5,
        color=COLORS["muted"],
        italic=True,
    )

    # One-slide design derivation: entropy component -> surprisal -> answer-level sum.
    c.rect(0.55, 1.52, 12.23, 2.16, COLORS["panel_header"], COLORS["panel_line"], 0.8, 0.09)

    c.rect(0.80, 1.80, 3.38, 1.55, COLORS["white"], COLORS["panel_line"], 0.7, 0.08)
    c.ellipse(1.00, 1.98, 0.31, 0.31, COLORS["title"])
    c.text("1", 1.00, 1.97, 0.31, 0.31, 12.0, color="FFFFFF", bold=True, align="center")
    c.text("Identify the strongest term", 1.42, 1.94, 2.43, 0.36, 14.5, bold=True)
    c.text("H_B(m₅)", 1.10, 2.36, 2.78, 0.43, 22.0, family="math", bold=True, align="center")
    c.rect(1.14, 2.79, 2.70, 0.27, COLORS["green_light"], radius=0.05)
    c.text("escape 0.742  >  full entropy 0.726", 1.15, 2.77, 2.68, 0.30, 10.8, color=COLORS["green_dark"], bold=True, family="sans", align="center")
    c.text("discard within-top-5 and tail-composition terms", 1.00, 3.10, 2.98, 0.18, 9.8, color=COLORS["muted"], family="sans", align="center")

    c.text("→", 4.21, 2.32, 0.36, 0.43, 23.0, color=COLORS["title"], bold=True, align="center")

    c.rect(4.59, 1.80, 3.38, 1.55, COLORS["white"], COLORS["panel_line"], 0.7, 0.08)
    c.ellipse(4.79, 1.98, 0.31, 0.31, COLORS["title"])
    c.text("2", 4.79, 1.97, 0.31, 0.31, 12.0, color="FFFFFF", bold=True, align="center")
    c.text("Design a monotonic score", 5.20, 1.94, 2.49, 0.36, 14.5, bold=True)
    c.text("H_B(m₅)  →  −log m₅⁽ᵗ⁾", 4.82, 2.34, 2.92, 0.43, 16.0, family="math", bold=True, align="center")
    c.rect(4.97, 2.78, 2.62, 0.24, COLORS["orange_light"], radius=0.05)
    c.text("H_B is not globally monotonic", 4.98, 2.76, 2.60, 0.27, 9.9, color=COLORS["orange"], family="sans", align="center")
    c.rect(4.97, 3.06, 2.62, 0.22, COLORS["green_light"], radius=0.05)
    c.text("less top-5 mass → larger score", 4.98, 3.04, 2.60, 0.25, 10.0, color=COLORS["green_dark"], bold=True, family="sans", align="center")

    c.text("→", 8.00, 2.32, 0.36, 0.43, 23.0, color=COLORS["title"], bold=True, align="center")

    c.rect(8.38, 1.80, 4.10, 1.55, COLORS["green_light"], COLORS["green"], 0.8, 0.08)
    c.ellipse(8.58, 1.98, 0.31, 0.31, COLORS["green_dark"])
    c.text("3", 8.58, 1.97, 0.31, 0.31, 12.0, color="FFFFFF", bold=True, align="center")
    c.text("Aggregate over the answer", 8.99, 1.94, 3.15, 0.36, 15.0, color=COLORS["green_dark"], bold=True)
    c.text(
        "tl_sum = ∑ₜ −log m₅⁽ᵗ⁾\n= −log ∏ₜ m₅⁽ᵗ⁾",
        8.73,
        2.28,
        3.44,
        0.70,
        18.3,
        color=COLORS["green_dark"],
        family="math",
        bold=True,
        align="center",
    )
    c.text("joint surprisal of staying inside top-5", 8.78, 3.06, 3.34, 0.18, 9.9, color=COLORS["green_dark"], family="sans", align="center")

    c.rect(0.55, 3.88, 7.35, 2.36, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)
    c.rect(8.15, 3.88, 4.63, 2.36, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)
    c.text("Why sum?", 0.86, 4.04, 2.10, 0.34, 17.0, bold=True)
    c.text("wrong answers often leak a little at several positions", 3.13, 4.10, 4.32, 0.22, 10.4, color=COLORS["muted"], family="sans", align="right")

    token_specs = [
        (0.88, "token t₁", 0.94),
        (2.54, "token t₂", 0.90),
        (4.20, "token t₃", 0.87),
    ]
    for x, label, mass in token_specs:
        c.rect(x, 4.47, 1.30, 0.92, COLORS["panel_header"], COLORS["panel_line"], 0.6, 0.07)
        c.text(label, x + 0.05, 4.53, 1.20, 0.25, 11.5, bold=True, family="sans", align="center")
        c.rect(x + 0.16, 4.83, 0.98, 0.19, COLORS["orange_light"], radius=0.03)
        c.rect(x + 0.16, 4.83, 0.98 * mass, 0.19, COLORS["green"], radius=0.03)
        c.text("small leak", x + 0.08, 5.08, 1.14, 0.19, 9.4, color=COLORS["muted"], family="sans", align="center")
    c.text("+", 2.24, 4.72, 0.26, 0.30, 18.0, color=COLORS["muted"], bold=True, align="center")
    c.text("+", 3.90, 4.72, 0.26, 0.30, 18.0, color=COLORS["muted"], bold=True, align="center")
    c.text("→", 5.55, 4.69, 0.38, 0.34, 20.0, color=COLORS["title"], bold=True, align="center")
    c.rect(5.98, 4.47, 1.48, 0.92, COLORS["green_light"], COLORS["green"], 0.8, 0.07)
    c.text("Σ", 6.34, 4.50, 0.76, 0.31, 20.0, color=COLORS["green_dark"], bold=True, family="math", align="center")
    c.text("tl_sum", 6.12, 4.87, 1.20, 0.27, 13.0, color=COLORS["green_dark"], bold=True, family="sans", align="center")

    c.rect(0.91, 5.61, 2.78, 0.35, COLORS["track"], radius=0.06)
    c.text("Worst token: keeps one position", 0.92, 5.59, 2.76, 0.38, 10.6, color=COLORS["muted"], family="sans", align="center")
    c.rect(3.97, 5.61, 3.43, 0.35, COLORS["green_light"], radius=0.06)
    c.text("tl_sum: all positions accumulate", 3.98, 5.59, 3.41, 0.38, 10.8, color=COLORS["green_dark"], bold=True, family="sans", align="center")

    c.text("Aggregation scan", 8.44, 4.04, 2.28, 0.32, 16.0, bold=True)
    c.text("macro AUROC ↑", 10.98, 4.30, 1.45, 0.22, 10.0, color=COLORS["muted"], family="sans", align="right")
    track_x, track_w = 9.70, 2.17
    aggregations = [
        ("tl_sum", 0.7493, COLORS["green"], 4.58, True),
        ("min_top5m", 0.7419, COLORS["blue"], 5.13, False),
        ("tl_mean", 0.7133, COLORS["gray_bar"], 5.68, False),
    ]
    for label, value, fill, y, is_best in aggregations:
        c.text(label, 8.43, y - 0.06, 1.12, 0.34, 11.7, bold=is_best, family="sans")
        c.rect(track_x, y, track_w, 0.30, COLORS["track"], radius=0.05)
        bar_w = max(0.06, (value - 0.70) / 0.05 * track_w)
        c.rect(track_x, y, bar_w, 0.30, fill, radius=0.05)
        c.text(f"{value:.4f}", 11.91, y - 0.05, 0.58, 0.36, 12.0, bold=is_best, family="sans", align="right")
    c.rect(8.58, 6.02, 3.74, 0.17, COLORS["green_light"], radius=0.05)
    c.text(
        "tl_sum vs. min_top5m: paired bootstrap p < 0.001",
        8.59,
        5.99,
        3.72,
        0.23,
        10.1,
        color=COLORS["green_dark"],
        bold=True,
        family="sans",
        align="center",
    )

    c.rect(0.72, 6.44, 11.82, 0.50, COLORS["highlight"], radius=0.07)
    c.text(
        "Takeaway: tl_sum is built from the top-5 escape signal identified by entropy decomposition—not an ad hoc formula.",
        0.92,
        6.45,
        11.40,
        0.46,
        15.7,
        color=COLORS["green_dark"],
        bold=True,
        align="center",
    )
    return c

def add_appendix_top5(prs: Presentation) -> SlideCanvas:
    c = SlideCanvas(prs)
    c.header("Why Top-5? Deterministic, Not Random", "Appendix", "A1", 28.0)
    c.text(
        "At each decoding step, rank the full-vocabulary softmax probabilities.",
        0.66,
        1.12,
        12.0,
        0.34,
        15.2,
        color=COLORS["muted"],
        italic=True,
    )

    c.rect(0.55, 1.55, 7.43, 4.70, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)
    c.rect(8.18, 1.55, 4.60, 4.70, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)

    c.text("What happens at token t?", 0.86, 1.76, 3.35, 0.36, 18.0, bold=True)
    c.text("deterministic at every step", 5.54, 1.82, 1.94, 0.22, 10.2, color=COLORS["green_dark"], family="sans", align="right")

    flow_cards = [
        (0.92, "logits at t", "one score / token"),
        (3.00, "softmax", "p_t(v) over |V|"),
        (5.08, "sort ↓", "largest to smallest"),
    ]
    for x, label, sub in flow_cards:
        c.rect(x, 2.26, 1.55, 0.70, COLORS["panel_header"], COLORS["panel_line"], 0.7, 0.07)
        c.text(label, x + 0.08, 2.33, 1.39, 0.28, 14.2, bold=True, family="sans", align="center")
        c.text(sub, x + 0.06, 2.65, 1.43, 0.18, 9.5, color=COLORS["muted"], family="sans", align="center")
    c.text("→", 2.55, 2.40, 0.35, 0.35, 19.0, color=COLORS["title"], bold=True, align="center")
    c.text("→", 4.63, 2.40, 0.35, 0.35, 19.0, color=COLORS["title"], bold=True, align="center")

    c.text("Illustrative ranked distribution", 0.92, 3.13, 2.82, 0.22, 10.3, color=COLORS["muted"], family="sans")
    bar_x, bar_y, total_w = 0.95, 3.48, 6.55
    segments = [
        (0.42, "147A5B", "1"),
        (0.27, COLORS["green"], "2"),
        (0.12, "5FB795", "3"),
        (0.08, "8DC8B2", "4"),
        (0.05, "B7DDCE", "5"),
        (0.06, COLORS["orange"], "6…|V|"),
    ]
    cursor = bar_x
    for mass, fill, label in segments:
        width = total_w * mass
        c.rect(cursor, bar_y, width, 0.42, fill)
        if width >= 0.36:
            c.text(label, cursor, bar_y + 0.04, width, 0.30, 10.0, color="FFFFFF" if mass > 0.08 else COLORS["ink"], bold=True, family="sans", align="center")
        cursor += width
    c.text("TOP-5 candidate set A_t", 1.18, 3.94, 5.42, 0.25, 11.2, color=COLORS["green_dark"], bold=True, family="sans", align="center")
    c.text("TAIL", 6.72, 3.94, 0.67, 0.25, 10.6, color=COLORS["orange"], bold=True, family="sans", align="center")

    c.rect(0.91, 4.34, 6.69, 0.62, COLORS["panel_header"], radius=0.07)
    c.text("mₜ = ∑ᵢ₌₁⁵ p₍ᵢ₎⁽ᵗ⁾", 1.08, 4.40, 2.78, 0.39, 17.0, family="math", bold=True, align="center")
    c.text("tail mass = 1 − mₜ", 4.25, 4.40, 3.05, 0.39, 16.0, family="math", bold=True, align="center")

    c.rect(0.91, 5.14, 6.69, 0.38, COLORS["green_light"], radius=0.06)
    c.text("No sampling: A_t is always the five largest probabilities under p_t.", 1.02, 5.12, 6.47, 0.42, 11.8, color=COLORS["green_dark"], bold=True, family="sans", align="center")
    c.rect(0.91, 5.65, 6.69, 0.36, COLORS["orange_light"], radius=0.06)
    c.text("Tail means ranks 6…|V|—not the bottom five tokens.", 1.02, 5.63, 6.47, 0.40, 11.5, color=COLORS["orange"], bold=True, family="sans", align="center")

    c.text("Why choose k = 5?", 8.47, 1.76, 2.73, 0.36, 18.0, bold=True)
    c.text("tl_sum@k = Σ_t −log m_{k,t} · macro AUROC", 8.48, 2.15, 3.93, 0.24, 9.9, color=COLORS["muted"], family="sans")
    track_x, track_w = 9.55, 2.19
    cutoff_rows = [
        ("k = 2", 0.7412, COLORS["gray_bar"], 2.72, False),
        ("k = 5", 0.7493, COLORS["green"], 3.48, True),
        ("k = 10", 0.7498, COLORS["blue"], 4.24, False),
    ]
    for label, value, fill, y, selected in cutoff_rows:
        c.text(label, 8.50, y - 0.04, 0.83, 0.32, 12.2, bold=selected, family="sans")
        c.rect(track_x, y, track_w, 0.31, COLORS["track"], radius=0.05)
        width = max(0.05, (value - 0.72) / 0.03 * track_w)
        c.rect(track_x, y, width, 0.31, fill, radius=0.05)
        c.text(f"{value:.4f}", 11.81, y - 0.04, 0.62, 0.34, 12.0, bold=selected, family="sans", align="right")
    c.rect(8.48, 4.79, 3.95, 0.31, COLORS["green_light"], radius=0.06)
    c.text("k=5 ≈ k=10 (p=.824); k=2 is lower (p=.022)", 8.49, 4.77, 3.93, 0.35, 9.9, color=COLORS["green_dark"], bold=True, family="sans", align="center")
    c.ellipse(8.54, 5.27, 0.13, 0.13, COLORS["green"])
    c.text("Allows several casing, alias, and subword variants.", 8.76, 5.19, 3.63, 0.27, 10.5, family="sans")
    c.ellipse(8.54, 5.62, 0.13, 0.13, COLORS["green"])
    c.text("Still keeps the candidate set narrow.", 8.76, 5.54, 3.63, 0.27, 10.5, family="sans")
    c.text("Post-hoc sensitivity on the same frozen 4×1,000 direct traces.", 8.48, 5.92, 3.94, 0.21, 9.1, color=COLORS["muted"], family="sans", italic=True, align="center")

    c.rect(0.72, 6.44, 11.82, 0.50, COLORS["highlight"], radius=0.07)
    c.text(
        "Takeaway: top-5 is deterministic; k=5 is a validated but model- and tokenizer-dependent hyperparameter.",
        0.92,
        6.45,
        11.40,
        0.46,
        15.4,
        color=COLORS["green_dark"],
        bold=True,
        align="center",
    )
    return c


def add_appendix_derivation(prs: Presentation) -> SlideCanvas:
    c = SlideCanvas(prs)
    c.header("Entropy Decomposition: Step-by-Step", "Appendix", "A2", 28.0)
    c.text(
        "Split the vocabulary into the top-5 set A and the tail B; drop token index t for readability.",
        0.66,
        1.12,
        12.0,
        0.34,
        14.7,
        color=COLORS["muted"],
        italic=True,
    )

    c.rect(0.55, 1.51, 12.23, 0.95, COLORS["panel_header"], COLORS["panel_line"], 0.8, 0.08)
    c.text("A = V₅(p),   B = V \\ A", 0.83, 1.68, 2.66, 0.36, 16.0, family="math", bold=True, align="center")
    c.text("m = ∑ᵥ∈A p(v)", 3.68, 1.68, 2.55, 0.36, 16.0, family="math", bold=True, align="center")
    c.text("p_top(v) = p(v)/m", 6.45, 1.68, 2.66, 0.36, 15.0, family="math", align="center")
    c.text("p_tail(v) = p(v)/(1−m)", 9.25, 1.68, 3.10, 0.36, 15.0, family="math", align="center")
    c.text("for v in A", 7.08, 2.05, 1.37, 0.19, 9.4, color=COLORS["muted"], family="sans", align="center")
    c.text("for v in B", 10.16, 2.05, 1.35, 0.19, 9.4, color=COLORS["muted"], family="sans", align="center")

    c.rect(0.55, 2.68, 8.02, 3.57, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)
    c.rect(8.80, 2.68, 3.98, 3.57, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)
    c.text("Algebraic identity", 0.85, 2.86, 2.75, 0.34, 17.5, bold=True)
    c.text("exact—not an approximation", 5.76, 2.93, 2.35, 0.21, 9.9, color=COLORS["green_dark"], family="sans", align="right")

    derivation_steps = [
        ("1", "H(p) = −∑ᵥ∈A p(v) log p(v)  −  ∑ᵥ∈B p(v) log p(v)", 3.27, 0.46, 14.0),
        ("2", "= −∑ᵥ∈A m p_top(v) log[m p_top(v)]\n  −∑ᵥ∈B (1−m) p_tail(v) log[(1−m) p_tail(v)]", 3.84, 0.69, 12.1),
        ("3", "= −m log m − (1−m) log(1−m)\n  + m H(p_top) + (1−m) H(p_tail)", 4.64, 0.67, 13.0),
    ]
    for badge, formula, y, height, size in derivation_steps:
        c.ellipse(0.86, y + 0.06, 0.28, 0.28, COLORS["title"])
        c.text(badge, 0.86, y + 0.045, 0.28, 0.28, 11.0, color="FFFFFF", bold=True, align="center")
        c.text(formula, 1.30, y, 6.83, height, size, family="math", align="left")

    c.rect(0.91, 5.48, 7.28, 0.50, COLORS["green_light"], COLORS["green"], 0.8, 0.07)
    c.text(
        "H(p) = H_B(m) + m H(p_top) + (1−m) H(p_tail)",
        1.07,
        5.50,
        6.96,
        0.43,
        16.8,
        color=COLORS["green_dark"],
        family="math",
        bold=True,
        align="center",
    )

    c.text("Empirical probes (conditional; approximate)", 9.08, 2.84, 3.40, 0.31, 12.7, bold=True)
    c.text("The exact identity weights the last two by m and 1−m.", 9.08, 3.12, 3.40, 0.18, 8.8, color=COLORS["muted"], family="sans", align="center")
    term_cards = [
        ("H_B(m)", "top-5 vs. tail membership", "0.742", COLORS["green"], 3.34),
        ("H(p_top)", "competition inside top-5", "0.712", COLORS["blue"], 4.08),
        ("H(p_tail)", "composition inside the tail", "0.502", COLORS["orange"], 4.82),
    ]
    for formula, meaning, value, fill, y in term_cards:
        c.rect(9.08, y, 3.40, 0.60, COLORS["panel_header"], COLORS["panel_line"], 0.6, 0.06)
        c.text(formula, 9.21, y + 0.06, 1.96, 0.28, 11.9, bold=True, family="math")
        c.text(meaning, 9.21, y + 0.34, 2.45, 0.18, 9.5, color=COLORS["muted"], family="sans")
        c.rect(11.75, y + 0.14, 0.59, 0.28, fill, radius=0.05)
        c.text(value, 11.76, y + 0.12, 0.57, 0.31, 10.6, color="FFFFFF", bold=True, family="sans", align="center")

    c.rect(9.08, 5.54, 3.40, 0.24, COLORS["orange_light"], radius=0.05)
    c.text("Algebraic decomposition ends here.", 9.09, 5.52, 3.38, 0.28, 9.7, color=COLORS["orange"], bold=True, family="sans", align="center")
    c.rect(9.08, 5.84, 3.40, 0.29, COLORS["green_light"], radius=0.05)
    c.text("Design: TLₜ = −log mₜ;  tl_sum = ∑ₜ TLₜ", 9.09, 5.82, 3.38, 0.32, 9.9, color=COLORS["green_dark"], bold=True, family="math", align="center")

    c.rect(0.72, 6.44, 11.82, 0.50, COLORS["highlight"], radius=0.07)
    c.text(
        "Takeaway: decomposition locates the escape signal; surprisal is a motivated scoring choice, not an equality.",
        0.92,
        6.45,
        11.40,
        0.46,
        15.4,
        color=COLORS["green_dark"],
        bold=True,
        align="center",
    )
    return c
def new_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W_IN)
    prs.slide_height = Inches(SLIDE_H_IN)
    return prs


def validate(prs: Presentation, expected_slides: int):
    assert len(prs.slides) == expected_slides
    for index, slide in enumerate(prs.slides, start=1):
        assert len(slide.shapes) >= 20, f"slide {index}: unexpectedly few shapes"
        for shape in slide.shapes:
            assert shape.left + shape.width > 0
            assert shape.top + shape.height > 0


def save_pdf(canvases: list[SlideCanvas], path: Path):
    rgb_images = [canvas.image.convert("RGB") for canvas in canvases]
    rgb_images[0].save(
        path,
        "PDF",
        save_all=True,
        append_images=rgb_images[1:],
        resolution=192.0,
        quality=95,
    )


def main():
    os.environ.setdefault("PYTHONUTF8", "1")

    main_prs = new_presentation()
    main_slides = [add_slide_1(main_prs), add_slide_2(main_prs)]
    validate(main_prs, 2)
    main_prs.save(PPTX_OUT)
    save_pdf(main_slides, PDF_OUT)
    for canvas, png_path in zip(main_slides, PNG_OUT):
        canvas.image.save(png_path, "PNG", optimize=True)

    appendix_prs = new_presentation()
    appendix_slides = [
        add_appendix_top5(appendix_prs),
        add_appendix_derivation(appendix_prs),
    ]
    validate(appendix_prs, 2)
    appendix_prs.save(APPENDIX_PPTX_OUT)
    save_pdf(appendix_slides, APPENDIX_PDF_OUT)
    for canvas, png_path in zip(appendix_slides, APPENDIX_PNG_OUT):
        canvas.image.save(png_path, "PNG", optimize=True)

    full_prs = new_presentation()
    full_slides = [
        add_slide_1(full_prs),
        add_slide_2(full_prs),
        add_appendix_top5(full_prs),
        add_appendix_derivation(full_prs),
    ]
    validate(full_prs, 4)
    full_prs.save(FULL_PPTX_OUT)
    save_pdf(full_slides, FULL_PDF_OUT)

    outputs = [
        PPTX_OUT,
        PDF_OUT,
        APPENDIX_PPTX_OUT,
        APPENDIX_PDF_OUT,
        FULL_PPTX_OUT,
        FULL_PDF_OUT,
        *PNG_OUT,
        *APPENDIX_PNG_OUT,
    ]
    for path in outputs:
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
