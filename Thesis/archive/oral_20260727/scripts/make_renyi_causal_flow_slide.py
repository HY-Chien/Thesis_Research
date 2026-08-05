#!/usr/bin/env python3
"""Build a standalone causal-flow slide for the Rényi-alpha motivation."""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, "/home/S113065528/.codex_tmp/python_pptx")
sys.path.insert(0, "/tmp/seek_pptx_lib")

from PIL import Image
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from make_tl_sum_2slides import (
    COLORS,
    PREVIEW_DPI,
    SlideCanvas,
    new_presentation,
    pil_color,
    save_pdf,
    validate,
)


OUT_DIR = Path(__file__).resolve().parent
PPTX_OUT = OUT_DIR / "renyi_alpha_causal_flow_slide.pptx"
PDF_OUT = OUT_DIR / "renyi_alpha_causal_flow_slide.pdf"
PNG_OUT = OUT_DIR / "renyi_alpha_causal_flow_slide.png"


# Match the blue-gray thesis-defense deck shown in the current presentation.
COLORS.update(
    {
        "background": "FFFFFF",
        "white": "FFFFFF",
        "title": "064B6B",
        "ink": "111827",
        "muted": "52606D",
        "line": "7B9AAE",
        "panel_line": "B6C8D8",
        "panel_header": "E9ECEF",
        "ice": "EDF3F7",
        "ice_deep": "FFFFFF",
        "blue": "1767D2",
        "blue_light": "B8D1EC",
        "green": "1E9C72",
        "green_dark": "137455",
        "green_light": "E5F4EE",
        "highlight": "E7F3EC",
        "track": "EEF2F5",
    }
)

NAVY = "0B3972"
PALE_BLUE = "F4F8FC"
ORANGE = "E49325"
TAIL_BLUE = "4C86C6"
GRAY_1 = "BAC4CA"
GRAY_2 = "7892A1"


def render_formula(tex: str, name: str) -> Path:
    """Render a tightly cropped transparent formula PNG with local TeX tools."""
    temp_dir = Path(tempfile.mkdtemp(prefix=f"renyi_{name}_"))
    tex_path = temp_dir / "formula.tex"
    tex_path.write_text(
        "\n".join(
            [
                r"\documentclass{article}",
                r"\usepackage{amsmath}",
                r"\pagestyle{empty}",
                r"\begin{document}",
                rf"\[ {tex} \]",
                r"\end{document}",
            ]
        ),
        encoding="utf-8",
    )
    subprocess.run(
        [
            "/home/S113065528/texlive/2026/bin/x86_64-linux/latex",
            "-interaction=nonstopmode",
            "-halt-on-error",
            f"-output-directory={temp_dir}",
            str(tex_path),
        ],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.STDOUT,
    )
    png_path = temp_dir / f"{name}.png"
    subprocess.run(
        [
            "/home/S113065528/texlive/2026/bin/x86_64-linux/dvipng",
            "-D",
            "320",
            "-T",
            "tight",
            "-bg",
            "Transparent",
            "-o",
            str(png_path),
            str(temp_dir / "formula.dvi"),
        ],
        check=True,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.STDOUT,
    )
    return png_path


def add_picture_fit(c: SlideCanvas, path: Path, x: float, y: float, w: float, h: float):
    """Fit an image in a box in both PPTX and the matching PNG preview."""
    with Image.open(path) as source:
        image = source.convert("RGBA")
    source_w, source_h = image.size
    ratio = min(w * PREVIEW_DPI / source_w, h * PREVIEW_DPI / source_h)
    target_w = max(1, int(round(source_w * ratio)))
    target_h = max(1, int(round(source_h * ratio)))
    actual_w = target_w / PREVIEW_DPI
    actual_h = target_h / PREVIEW_DPI
    actual_x = x + (w - actual_w) / 2
    actual_y = y + (h - actual_h) / 2
    c.slide.shapes.add_picture(
        str(path),
        Inches(actual_x),
        Inches(actual_y),
        width=Inches(actual_w),
        height=Inches(actual_h),
    )
    resized = image.resize((target_w, target_h), Image.Resampling.LANCZOS)
    c.image.paste(
        resized,
        (c.px(actual_x), c.px(actual_y)),
        resized,
    )


def add_right_arrow(c: SlideCanvas, x: float, y: float, w: float, h: float):
    shape = c.slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = c.slide.background.fill.fore_color.rgb
    shape.fill.fore_color.rgb = shape.line.color.rgb = __import__(
        "make_tl_sum_2slides"
    ).rgb(COLORS["blue"])
    shape.line.fill.background()

    x0, y0 = c.px(x), c.px(y)
    wp, hp = c.px(w), c.px(h)
    c.draw.polygon(
        [
            (x0, y0 + int(hp * 0.27)),
            (x0 + int(wp * 0.56), y0 + int(hp * 0.27)),
            (x0 + int(wp * 0.56), y0),
            (x0 + wp, y0 + hp // 2),
            (x0 + int(wp * 0.56), y0 + hp),
            (x0 + int(wp * 0.56), y0 + int(hp * 0.73)),
            (x0, y0 + int(hp * 0.73)),
        ],
        fill=pil_color(COLORS["blue"]),
    )


def add_panel(c: SlideCanvas, x: float, y: float, w: float, h: float):
    c.rect(x + 0.04, y + 0.05, w, h, "E1E6EA", radius=0.10)
    c.rect(x, y, w, h, COLORS["white"], COLORS["panel_line"], 0.8, 0.10)


def add_step_header(
    c: SlideCanvas,
    x: float,
    w: float,
    number: str,
    verb: str,
    descriptor: str,
):
    c.ellipse(x + 0.20, 1.82, 0.42, 0.42, NAVY)
    c.text(
        number,
        x + 0.20,
        1.82,
        0.42,
        0.42,
        12.5,
        color=COLORS["white"],
        bold=True,
        family="sans",
        align="center",
        margin=0,
    )
    c.text(
        verb,
        x + 0.74,
        1.78,
        w - 0.94,
        0.27,
        13.3,
        color=NAVY,
        bold=True,
        family="sans",
    )
    c.text(
        descriptor,
        x + 0.74,
        2.05,
        w - 0.94,
        0.24,
        10.2,
        color=COLORS["muted"],
        family="sans",
    )


def add_formula_panel(c: SlideCanvas, x: float, y: float, w: float, h: float):
    add_panel(c, x, y, w, h)
    add_step_header(c, x, w, "1", "DEFINE", "Connect Rényi to Shannon")

    renyi_formula = render_formula(
        r"\displaystyle H_{\alpha}(p)=\frac{1}{1-\alpha}"
        r"\log\sum_{v}p(v)^{\alpha}",
        "renyi",
    )
    shannon_formula = render_formula(
        r"\displaystyle \lim_{\alpha\to1}H_{\alpha}(p)"
        r"=H(p)=-\sum_{v}p(v)\log p(v)",
        "shannon",
    )
    c.text(
        "Rényi entropy family",
        x + 0.20,
        2.46,
        w - 0.40,
        0.28,
        12.8,
        color=COLORS["muted"],
        bold=True,
        family="sans",
        align="center",
    )
    add_picture_fit(c, renyi_formula, x + 0.18, 2.79, w - 0.36, 0.72)

    c.rect(x + 1.00, 3.65, w - 2.00, 0.40, COLORS["ice"], radius=0.06)
    c.text(
        "α → 1",
        x + 1.00,
        3.65,
        w - 2.00,
        0.40,
        13.5,
        color=NAVY,
        bold=True,
        family="math",
        align="center",
        margin=0,
    )
    add_picture_fit(c, shannon_formula, x + 0.10, 4.23, w - 0.20, 0.72)

    c.rect(x + 0.26, 5.22, w - 0.52, 0.55, PALE_BLUE, COLORS["blue_light"], 0.7, 0.07)
    c.text(
        "α → 1 recovers Shannon entropy.",
        x + 0.34,
        5.21,
        w - 0.68,
        0.57,
        10.5,
        color=NAVY,
        bold=True,
        family="sans",
        align="center",
    )
    c.text(
        "One parameter controls which probabilities matter.",
        x + 0.28,
        5.90,
        w - 0.56,
        0.24,
        9.8,
        color=COLORS["muted"],
        italic=True,
        family="sans",
        align="center",
    )


def add_weighting_panel(c: SlideCanvas, x: float, y: float, w: float, h: float):
    add_panel(c, x, y, w, h)
    add_step_header(c, x, w, "2", "REWEIGHT", "See what α changes")
    c.text(
        "Same probabilities, different emphasis",
        x + 0.24,
        2.40,
        w - 0.48,
        0.30,
        12.7,
        color=NAVY,
        bold=True,
        family="sans",
        align="center",
    )
    c.text(
        "Effective weight  ∝  p(v)^α",
        x + 0.24,
        2.72,
        w - 0.48,
        0.27,
        11.0,
        color=COLORS["muted"],
        family="math",
        align="center",
    )

    probabilities = [
        0.32,
        0.21,
        0.14,
        0.09,
        0.07,
        0.05,
        0.035,
        0.027,
        0.021,
        0.017,
        0.012,
        0.008,
    ]
    mini_panels = [
        (x + 0.23, 2.0, "α = 2", "HEAD-SENSITIVE"),
        (x + 1.92, 1.0, "α → 1", "SHANNON"),
        (x + 3.61, 0.5, "α = 0.5", "TAIL-SENSITIVE"),
    ]
    mini_w = 1.45
    mini_y = 3.10
    mini_h = 2.48
    base_y = 5.08
    max_height = 1.34
    common_y_max = 0.60
    bar_w = 0.055
    gap = 0.025

    for panel_x, alpha, alpha_label, interpretation in mini_panels:
        fill = COLORS["green_light"] if alpha < 1 else "F7F9FA"
        border = COLORS["green"] if alpha < 1 else COLORS["panel_line"]
        c.rect(panel_x, mini_y, mini_w, mini_h, fill, border, 0.65, 0.07)
        c.text(
            alpha_label,
            panel_x + 0.06,
            3.23,
            mini_w - 0.12,
            0.26,
            12.0,
            color=NAVY,
            bold=True,
            family="sans",
            align="center",
        )
        c.text(
            interpretation,
            panel_x + 0.04,
            3.50,
            mini_w - 0.08,
            0.22,
            7.9,
            color=COLORS["green_dark"] if alpha < 1 else COLORS["muted"],
            bold=alpha < 1,
            family="sans",
            align="center",
        )
        transformed = [value**alpha for value in probabilities]
        total = sum(transformed)
        weights = [value / total for value in transformed]
        c.line(panel_x + 0.15, base_y, mini_w - 0.30, COLORS["panel_line"], 0.008)
        bar_x = panel_x + 0.17
        for index, weight in enumerate(weights):
            if index == 5:
                bar_x += 0.05
            bar_h = weight / common_y_max * max_height
            c.rect(
                bar_x,
                base_y - bar_h,
                bar_w,
                bar_h,
                ORANGE if index < 5 else TAIL_BLUE,
            )
            bar_x += bar_w + gap
        c.text(
            "head",
            panel_x + 0.12,
            5.18,
            0.58,
            0.19,
            7.5,
            color="C97812",
            bold=True,
            family="sans",
            align="center",
        )
        c.text(
            "tail",
            panel_x + 0.83,
            5.18,
            0.48,
            0.19,
            7.5,
            color="3D75B3",
            bold=True,
            family="sans",
            align="center",
        )

    c.rect(x + 0.55, 5.83, w - 1.10, 0.44, PALE_BLUE, radius=0.07)
    c.text(
        "Lower α   →   stronger tail influence",
        x + 0.60,
        5.82,
        w - 1.20,
        0.46,
        12.0,
        color=NAVY,
        bold=True,
        family="sans",
        align="center",
    )


def add_result_panel(c: SlideCanvas, x: float, y: float, w: float, h: float):
    add_panel(c, x, y, w, h)
    add_step_header(c, x, w, "3", "DETECT", "Measure the empirical effect")
    c.text(
        "2Wiki error-detection AUROC ↑",
        x + 0.18,
        2.42,
        w - 0.36,
        0.28,
        11.7,
        color=NAVY,
        bold=True,
        family="sans",
        align="center",
    )
    c.rect(x + 1.01, 2.80, w - 2.02, 0.42, COLORS["green_light"], radius=0.07)
    c.text(
        "+0.065",
        x + 1.01,
        2.80,
        w - 2.02,
        0.42,
        12.7,
        color=COLORS["green_dark"],
        bold=True,
        family="sans",
        align="center",
        margin=0,
    )

    baseline_y = 5.35
    chart_top = 3.35
    chart_height = baseline_y - chart_top
    c.line(x + 0.34, baseline_y, w - 0.68, COLORS["panel_line"], 0.012)
    c.text(
        "0.50",
        x + 0.04,
        baseline_y - 0.12,
        0.35,
        0.18,
        7.8,
        color=COLORS["muted"],
        family="sans",
        align="right",
    )

    bars = [
        (x + 0.54, 0.651, GRAY_1, "α = 2"),
        (x + 1.34, 0.671, GRAY_2, "α → 1"),
        (x + 2.14, 0.716, COLORS["green"], "α = 0.5"),
    ]
    bar_w = 0.50
    for bar_x, value, fill, label in bars:
        bar_h = (value - 0.50) / 0.23 * chart_height
        bar_y = baseline_y - bar_h
        c.rect(bar_x, bar_y, bar_w, bar_h, fill, radius=0.06)
        c.text(
            f"{value:.3f}",
            bar_x - 0.12,
            bar_y - 0.35,
            bar_w + 0.24,
            0.28,
            11.5,
            bold=True,
            family="sans",
            align="center",
        )
        c.text(
            label,
            bar_x - 0.16,
            5.48,
            bar_w + 0.32,
            0.27,
            9.4,
            color=NAVY,
            bold=True,
            family="sans",
            align="center",
        )

    c.text(
        "Tail-sensitive α = 0.5 performs best.",
        x + 0.24,
        5.93,
        w - 0.48,
        0.27,
        10.0,
        color=COLORS["green_dark"],
        bold=True,
        family="sans",
        align="center",
    )


def add_slide(prs) -> SlideCanvas:
    c = SlideCanvas(prs)
    c.header(
        "Lower α Reveals Tail Uncertainty",
        "Methodology",
        "",
        29.0,
    )
    c.text(
        "Formula relationship  →  weighting mechanism  →  empirical error detection",
        0.66,
        1.12,
        12.0,
        0.30,
        13.0,
        color=COLORS["muted"],
        italic=True,
        family="sans",
        align="center",
    )

    panel_y, panel_h = 1.62, 4.80
    add_formula_panel(c, 0.62, panel_y, 2.76, panel_h)
    add_right_arrow(c, 3.48, 3.67, 0.30, 0.38)
    add_weighting_panel(c, 3.88, panel_y, 5.22, panel_h)
    add_right_arrow(c, 9.20, 3.67, 0.30, 0.38)
    add_result_panel(c, 9.60, panel_y, 3.10, panel_h)

    c.rect(0.92, 6.60, 11.50, 0.42, COLORS["highlight"], radius=0.07)
    c.text(
        "As α decreases, tail influence and error-detection AUROC both increase.",
        1.06,
        6.59,
        11.22,
        0.44,
        14.3,
        color=COLORS["green_dark"],
        bold=True,
        align="center",
    )
    return c


def main() -> None:
    prs = new_presentation()
    canvas = add_slide(prs)
    validate(prs, 1)
    prs.save(PPTX_OUT)
    save_pdf([canvas], PDF_OUT)
    canvas.image.save(PNG_OUT, "PNG", optimize=True)
    for path in (PPTX_OUT, PDF_OUT, PNG_OUT):
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
