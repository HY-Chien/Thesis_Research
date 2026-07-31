#!/usr/bin/env python3
"""Build an intuitive, editable slide for path-aware title retrieval."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, "/home/S113065528/.codex_tmp/python_pptx")
sys.path.insert(0, "/tmp/seek_pptx_lib")

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from make_tl_sum_2slides import (
    COLORS,
    SlideCanvas,
    new_presentation,
    pil_color,
    rgb,
    save_pdf,
    validate,
)


OUT_DIR = Path(__file__).resolve().parent
PPTX_OUT = OUT_DIR / "pathaware_title_retrieval_visual.pptx"
PDF_OUT = OUT_DIR / "pathaware_title_retrieval_visual.pdf"
PNG_OUT = OUT_DIR / "pathaware_title_retrieval_visual.png"


# Match the blue-gray thesis-defense deck.
COLORS.update(
    {
        "background": "FFFFFF",
        "white": "FFFFFF",
        "title": "064B6B",
        "ink": "111827",
        "muted": "52606D",
        "line": "7B9AAE",
        "panel_line": "B6C8D8",
        "panel_header": "E9ECEF",
        "ice": "EDF3F7",
        "ice_deep": "FFFFFF",
        "blue": "1767D2",
        "blue_light": "DCE9F7",
        "green": "1E9C72",
        "green_dark": "137455",
        "green_light": "E5F4EE",
        "highlight": "E7F3EC",
        "orange": "C97812",
        "orange_light": "FCEBD4",
        "red": "B64B4B",
        "red_light": "F9E2E2",
        "track": "EEF2F5",
    }
)

NAVY = "0B3972"
PALE_BLUE = "F4F8FC"


def add_arrow(c: SlideCanvas, x: float, y: float, w: float, h: float, color: str):
    shape = c.slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.line.fill.background()

    x0, y0 = c.px(x), c.px(y)
    wp, hp = c.px(w), c.px(h)
    c.draw.polygon(
        [
            (x0, y0 + int(hp * 0.28)),
            (x0 + int(wp * 0.58), y0 + int(hp * 0.28)),
            (x0 + int(wp * 0.58), y0),
            (x0 + wp, y0 + hp // 2),
            (x0 + int(wp * 0.58), y0 + hp),
            (x0 + int(wp * 0.58), y0 + int(hp * 0.72)),
            (x0, y0 + int(hp * 0.72)),
        ],
        fill=pil_color(color),
    )


def add_query_lines(c: SlideCanvas, x: float, y: float, widths: list[float], color: str):
    for index, width in enumerate(widths):
        c.rect(x, y + index * 0.17, width, 0.075, color, radius=0.035)


def add_switch(
    c: SlideCanvas,
    x: float,
    y: float,
    enabled: bool,
    label: str,
    detail: str,
):
    fill = COLORS["green_light"] if enabled else COLORS["blue_light"]
    accent = COLORS["green"] if enabled else COLORS["blue"]
    c.rect(x, y, 2.02, 1.16, fill, accent, 1.0, 0.10)
    c.text(
        "TITLE MATCHING",
        x + 0.12,
        y + 0.11,
        1.25,
        0.22,
        9.0,
        color=COLORS["muted"],
        bold=True,
        family="sans",
    )
    c.rect(x + 1.39, y + 0.10, 0.50, 0.27, accent, radius=0.13)
    knob_x = x + (1.65 if enabled else 1.42)
    c.ellipse(knob_x, y + 0.125, 0.21, 0.21, COLORS["white"])
    c.text(
        "ON" if enabled else "OFF",
        x + 0.12,
        y + 0.38,
        1.76,
        0.30,
        15.0,
        color=accent,
        bold=True,
        family="sans",
    )
    c.text(
        label,
        x + 0.12,
        y + 0.71,
        1.76,
        0.21,
        10.0,
        color=NAVY,
        bold=True,
        family="sans",
    )
    c.text(
        detail,
        x + 0.12,
        y + 0.92,
        1.76,
        0.17,
        8.6,
        color=COLORS["muted"],
        family="sans",
    )


def add_path_row(
    c: SlideCanvas,
    y: float,
    direct: bool,
):
    accent = COLORS["blue"] if direct else COLORS["green"]
    light = PALE_BLUE if direct else "F2FAF6"
    c.rect(0.66, y, 8.33, 1.83, light, COLORS["panel_line"], 0.8, 0.11)

    badge_text = "DIRECT PATH" if direct else "DECOMPOSITION PATH"
    badge_w = 1.52 if direct else 2.16
    c.rect(0.86, y + 0.16, badge_w, 0.31, accent, radius=0.07)
    c.text(
        badge_text,
        0.86,
        y + 0.155,
        badge_w,
        0.32,
        9.3,
        color=COLORS["white"],
        bold=True,
        family="sans",
        align="center",
        margin=0,
    )

    if direct:
        c.text(
            "Long original question",
            0.88,
            y + 0.55,
            2.72,
            0.27,
            14.2,
            color=NAVY,
            bold=True,
            family="sans",
        )
        c.text(
            "More words  ·  multiple constraints",
            0.88,
            y + 0.85,
            2.72,
            0.22,
            9.8,
            color=COLORS["muted"],
            family="sans",
        )
        add_query_lines(
            c,
            0.90,
            y + 1.22,
            [2.48, 2.18, 1.42],
            "86A9C9",
        )
        switch_label = "Passage content only"
        switch_detail = "Keep the full question"
        result_title = "Avoid noisy title hits"
        result_detail = "A stray keyword can retrieve\nan irrelevant article."
        result_fill = COLORS["red_light"]
        result_color = COLORS["red"]
    else:
        c.text(
            "Short sub-question",
            0.88,
            y + 0.55,
            2.72,
            0.27,
            14.2,
            color=NAVY,
            bold=True,
            family="sans",
        )
        c.text(
            "Entity-focused  ·  one knowledge need",
            0.88,
            y + 0.85,
            2.72,
            0.22,
            9.8,
            color=COLORS["muted"],
            family="sans",
        )
        c.rect(0.90, y + 1.18, 2.35, 0.39, COLORS["white"], accent, 0.8, 0.07)
        c.text(
            "Who directed Inception?",
            1.00,
            y + 1.18,
            2.16,
            0.39,
            11.0,
            color=NAVY,
            bold=True,
            family="sans",
            align="center",
        )
        switch_label = "Title + passage"
        switch_detail = "Use the entity as an anchor"
        result_title = "Find the right article"
        result_detail = "The title points directly to\nthe target entity."
        result_fill = COLORS["green_light"]
        result_color = COLORS["green_dark"]

    add_arrow(c, 3.58, y + 0.73, 0.42, 0.36, accent)
    add_switch(
        c,
        4.16,
        y + 0.34,
        enabled=not direct,
        label=switch_label,
        detail=switch_detail,
    )
    add_arrow(c, 6.35, y + 0.73, 0.42, 0.36, accent)
    c.rect(6.92, y + 0.34, 1.85, 1.16, result_fill, result_color, 1.0, 0.10)
    symbol = "OK" if not direct else "×"
    c.ellipse(7.60, y + 0.47, 0.46, 0.46, result_color)
    c.text(
        symbol,
        7.60,
        y + 0.455,
        0.46,
        0.46,
        10.5 if not direct else 18.0,
        color=COLORS["white"],
        bold=True,
        family="sans",
        align="center",
        margin=0,
    )
    c.text(
        result_title,
        7.05,
        y + 0.94,
        1.58,
        0.23,
        10.4,
        color=result_color,
        bold=True,
        family="sans",
        align="center",
    )
    c.text(
        result_detail,
        7.04,
        y + 1.18,
        1.60,
        0.29,
        7.9,
        color=COLORS["muted"],
        family="sans",
        align="center",
    )


def add_results_panel(c: SlideCanvas):
    x, y, w, h = 9.25, 1.55, 3.46, 4.76
    c.rect(x + 0.04, y + 0.05, w, h, "E1E6EA", radius=0.11)
    c.rect(x, y, w, h, COLORS["white"], COLORS["panel_line"], 0.9, 0.11)
    c.rect(x + 0.22, y + 0.20, w - 0.44, 0.46, COLORS["panel_header"], radius=0.08)
    c.text(
        "TITLE POLICY COMPARISON",
        x + 0.25,
        y + 0.19,
        w - 0.50,
        0.47,
        12.0,
        color=NAVY,
        bold=True,
        family="sans",
        align="center",
    )

    policies = [
        ("UNIFORM OFF", False, False, "F4F6F7", "BAC4CA"),
        ("UNIFORM ON", True, True, PALE_BLUE, "7892A1"),
        ("PATH-AWARE", False, True, "F2FAF6", COLORS["green"]),
    ]
    for index, (name, direct_on, decomp_on, card_fill, card_line) in enumerate(policies):
        card_x = x + 0.23 + index * 1.01
        c.rect(card_x, y + 0.80, 0.92, 1.02, card_fill, card_line, 0.9, 0.07)
        c.text(
            name,
            card_x + 0.03,
            y + 0.84,
            0.86,
            0.20,
            7.5,
            color=COLORS["green_dark"] if name == "PATH-AWARE" else NAVY,
            bold=True,
            family="sans",
            align="center",
        )
        for row, (path_name, enabled) in enumerate(
            [("Direct", direct_on), ("Decomp", decomp_on)]
        ):
            pill_y = y + 1.10 + row * 0.31
            pill_fill = COLORS["green_light"] if enabled else "E6EBEF"
            pill_color = COLORS["green_dark"] if enabled else COLORS["muted"]
            c.rect(card_x + 0.08, pill_y, 0.76, 0.25, pill_fill, radius=0.05)
            c.text(
                f"{path_name}  {'ON' if enabled else 'OFF'}",
                card_x + 0.09,
                pill_y,
                0.74,
                0.25,
                7.5,
                color=pill_color,
                bold=True,
                family="sans",
                align="center",
                margin=0,
            )

    c.text(
        "Macro-average EM",
        x + 0.26,
        y + 1.98,
        w - 0.52,
        0.25,
        10.0,
        color=COLORS["muted"],
        family="sans",
        align="center",
    )

    baseline_y = y + 4.14
    c.line(x + 0.34, baseline_y, w - 0.68, COLORS["panel_line"], 0.012)
    bars = [
        (x + 0.41, 39.3, 1.05, "BAC4CA"),
        (x + 1.42, 40.9, 1.42, "7892A1"),
        (x + 2.43, 41.8, 1.70, COLORS["green"]),
    ]
    for bar_x, value, bar_h, fill in bars:
        bar_y = baseline_y - bar_h
        c.rect(bar_x, bar_y, 0.56, bar_h, fill, radius=0.05)
        c.text(
            f"{value:.1f}",
            bar_x - 0.13,
            bar_y - 0.34,
            0.82,
            0.27,
            12.5,
            color=COLORS["green_dark"] if fill == COLORS["green"] else COLORS["ink"],
            bold=True,
            family="sans",
            align="center",
        )

    c.rect(x + 0.38, y + 4.31, w - 0.76, 0.29, COLORS["green_light"], radius=0.06)
    c.text(
        "Path-aware: best EM on 4 / 4 datasets",
        x + 0.40,
        y + 4.30,
        w - 0.80,
        0.31,
        9.1,
        color=COLORS["green_dark"],
        bold=True,
        family="sans",
        align="center",
    )


def add_slide(prs) -> SlideCanvas:
    c = SlideCanvas(prs)
    c.header(
        "Path-aware Retrieval Matches the Query",
        "Experiment",
        "48",
        27.5,
    )
    c.text(
        "Same retriever, different query forms  →  use a path-specific title policy",
        0.66,
        1.11,
        12.0,
        0.31,
        13.0,
        color=COLORS["muted"],
        italic=True,
        family="sans",
        align="center",
    )

    add_path_row(c, 1.55, direct=True)
    add_path_row(c, 3.73, direct=False)
    add_results_panel(c)

    c.rect(0.82, 6.42, 11.66, 0.51, COLORS["highlight"], radius=0.08)
    c.text(
        "Adaptive retrieval is not only whether to retrieve — it is also how to retrieve for each path.",
        0.98,
        6.42,
        11.34,
        0.51,
        14.0,
        color=COLORS["green_dark"],
        bold=True,
        family="serif",
        align="center",
    )
    return c


def main() -> None:
    prs = new_presentation()
    canvas = add_slide(prs)
    validate(prs, 1)
    prs.save(PPTX_OUT)
    save_pdf([canvas], PDF_OUT)
    canvas.image.save(PNG_OUT, "PNG", optimize=True)
    for path in (PPTX_OUT, PDF_OUT, PNG_OUT):
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
