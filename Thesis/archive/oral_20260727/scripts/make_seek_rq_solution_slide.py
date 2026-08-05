#!/usr/bin/env python3
"""Create a one-slide SEEK motivation deck in the thesis presentation style."""

from __future__ import annotations

import sys
from pathlib import Path

# The dependency is installed in /tmp for this workspace session.
sys.path.insert(0, "/tmp/seek_pptx_lib")

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "SEEK_Motivation_Research_Question_Solution.pptx"
PNG_PATH = OUT_DIR / "SEEK_Motivation_Research_Question_Solution_preview.png"

SLIDE_W = 13.333
SLIDE_H = 7.5

TEAL = "064B6B"
NAVY = "0B3972"
BLUE = "1767D2"
TEXT = "111827"
MUTED = "52606D"
LINE = "7B9AAE"
LIGHT_BORDER = "B6C8D8"
LIGHT_BLUE = "EDF5FF"
PALE_BLUE = "F6FAFE"
PALE_GRAY = "F3F5F7"
TAB_GRAY = "E9ECEF"
SHADOW = "E1E6EA"
WHITE = "FFFFFF"
CORNER = "AFC1CD"

SERIF = "Georgia"
SANS = "Arial"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_shape_fill(shape, color: str) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)


def set_shape_line(shape, color: str | None, width: float = 1.0) -> None:
    if color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(color)
        shape.line.width = Pt(width)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float,
    color: str = TEXT,
    font: str = SANS,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.MIDDLE,
    margin: float = 0.03,
    italic: bool = False,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign

    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1.0

    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return box


def add_rounded_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str,
    line: str | None = LIGHT_BORDER,
    line_width: float = 1.0,
    radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE,
):
    shape = slide.shapes.add_shape(
        radius_shape, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    set_shape_fill(shape, fill)
    set_shape_line(shape, line, line_width)
    return shape


def add_panel(slide, x: float, y: float, w: float, h: float):
    add_rounded_rect(
        slide,
        x + 0.045,
        y + 0.055,
        w,
        h,
        fill=SHADOW,
        line=None,
    )
    return add_rounded_rect(
        slide,
        x,
        y,
        w,
        h,
        fill=WHITE,
        line=LIGHT_BORDER,
        line_width=1.0,
    )


def add_circle_badge(
    slide,
    label: str,
    x: float,
    y: float,
    d: float,
    *,
    fill: str = NAVY,
    line: str | None = None,
    text_color: str = WHITE,
    size: float = 15,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)
    )
    set_shape_fill(shape, fill)
    set_shape_line(shape, line, 1.2)
    add_text(
        slide,
        label,
        x,
        y,
        d,
        d,
        size=size,
        color=text_color,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )
    return shape


def add_connector(slide, x1: float, y1: float, x2: float, y2: float, *, width=1.4):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = rgb(LINE)
    line.line.width = Pt(width)
    return line


def add_flow_card(
    slide,
    title: str,
    subtitle: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    icon: str,
):
    add_rounded_rect(slide, x, y, w, h, fill=PALE_BLUE, line=LIGHT_BORDER)
    add_circle_badge(
        slide,
        icon,
        x + 0.13,
        y + 0.19,
        0.44,
        fill=WHITE,
        line=BLUE,
        text_color=BLUE,
        size=14,
    )
    add_text(
        slide,
        title,
        x + 0.67,
        y + 0.12,
        w - 0.79,
        0.34,
        size=13.5,
        color=NAVY,
        bold=True,
        valign=MSO_ANCHOR.BOTTOM,
    )
    add_text(
        slide,
        subtitle,
        x + 0.67,
        y + 0.46,
        w - 0.79,
        h - 0.53,
        size=10.8,
        color=MUTED,
        valign=MSO_ANCHOR.TOP,
    )


def build_pptx() -> None:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(WHITE)

    # Original presentation header.
    add_text(
        slide,
        "Motivation",
        0.40,
        0.16,
        2.15,
        0.52,
        size=28,
        color=TEAL,
        font=SERIF,
        bold=True,
        margin=0,
    )
    top_rule = add_connector(slide, 1.02, 0.82, 12.30, 0.82, width=1.0)
    top_rule.line.color.rgb = rgb(LINE)

    tab = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(10.86),
        Inches(0.00),
        Inches(2.47),
        Inches(0.62),
    )
    set_shape_fill(tab, TAB_GRAY)
    set_shape_line(tab, None)
    add_text(
        slide,
        "Introduction",
        10.96,
        0.10,
        2.24,
        0.40,
        size=18,
        color="000000",
        font=SERIF,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    # Two-column content panels.
    left_x, panel_y, panel_w, panel_h = 0.68, 1.20, 5.55, 5.15
    right_x = 7.10
    add_panel(slide, left_x, panel_y, panel_w, panel_h)
    add_panel(slide, right_x, panel_y, panel_w, panel_h)

    # Middle transition arrow.
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(6.38),
        Inches(3.45),
        Inches(0.45),
        Inches(0.43),
    )
    set_shape_fill(arrow, BLUE)
    set_shape_line(arrow, None)

    # Panel headings.
    add_circle_badge(slide, "1", 0.98, 1.47, 0.53, size=15)
    add_text(
        slide,
        "RESEARCH QUESTION",
        1.67,
        1.43,
        3.95,
        0.55,
        size=18,
        color=NAVY,
        font=SERIF,
        bold=True,
        margin=0,
    )
    add_circle_badge(slide, "2", 7.40, 1.47, 0.53, size=15)
    add_text(
        slide,
        "SOLUTION — SEEK",
        8.09,
        1.43,
        3.95,
        0.55,
        size=18,
        color=NAVY,
        font=SERIF,
        bold=True,
        margin=0,
    )

    # Left panel: decomposition benefit versus risk.
    decomposition = add_rounded_rect(
        slide,
        2.30,
        2.14,
        2.28,
        0.55,
        fill=NAVY,
        line=None,
    )
    add_text(
        slide,
        "DECOMPOSITION",
        2.30,
        2.14,
        2.28,
        0.55,
        size=14,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    add_connector(slide, 3.44, 2.69, 3.44, 2.90)
    add_connector(slide, 2.00, 2.90, 4.88, 2.90)
    add_connector(slide, 2.00, 2.90, 2.00, 3.06)
    add_connector(slide, 4.88, 2.90, 4.88, 3.06)

    add_rounded_rect(
        slide, 1.12, 3.06, 2.35, 0.96, fill=LIGHT_BLUE, line="9FBCE0"
    )
    add_circle_badge(
        slide,
        "+",
        1.28,
        3.27,
        0.46,
        fill=WHITE,
        line=BLUE,
        text_color=BLUE,
        size=17,
    )
    add_text(
        slide,
        "Focused\nretrieval",
        1.89,
        3.15,
        1.42,
        0.75,
        size=14,
        color=NAVY,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    add_rounded_rect(
        slide, 3.62, 3.06, 2.14, 0.96, fill=PALE_GRAY, line="BCC6CE"
    )
    add_circle_badge(
        slide,
        "!",
        3.77,
        3.27,
        0.46,
        fill=WHITE,
        line=NAVY,
        text_color=NAVY,
        size=16,
    )
    add_text(
        slide,
        "Context loss\nError propagation",
        4.33,
        3.14,
        1.29,
        0.76,
        size=10.5,
        color=TEXT,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    # Main research question.
    add_rounded_rect(
        slide,
        1.10,
        4.42,
        4.72,
        1.30,
        fill=PALE_BLUE,
        line=BLUE,
        line_width=1.3,
    )
    add_text(
        slide,
        "How can adaptive RAG benefit from\ndecomposition without trusting it blindly?",
        1.34,
        4.53,
        4.24,
        1.05,
        size=16.5,
        color=TEXT,
        font=SERIF,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0.02,
    )

    # Right panel: concise SEEK dual-path solution.
    add_rounded_rect(
        slide,
        8.60,
        2.13,
        2.55,
        0.54,
        fill=NAVY,
        line=None,
    )
    add_text(
        slide,
        "ORIGINAL QUESTION",
        8.60,
        2.13,
        2.55,
        0.54,
        size=13.5,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    add_connector(slide, 9.88, 2.67, 9.88, 2.89)
    add_connector(slide, 8.50, 2.89, 11.25, 2.89)
    add_connector(slide, 8.50, 2.89, 8.50, 3.06)
    add_connector(slide, 11.25, 2.89, 11.25, 3.06)

    add_flow_card(
        slide,
        "Direct path",
        "Preserve context",
        7.43,
        3.06,
        2.36,
        1.04,
        icon="↔",
    )
    add_flow_card(
        slide,
        "Decomposed path",
        "Retrieve if uncertain",
        9.98,
        3.06,
        2.36,
        1.04,
        icon="?",
    )

    add_connector(slide, 8.61, 4.10, 8.61, 4.35)
    add_connector(slide, 11.16, 4.10, 11.16, 4.35)
    add_connector(slide, 8.61, 4.35, 11.16, 4.35)
    add_connector(slide, 9.88, 4.35, 9.88, 4.58)

    add_rounded_rect(
        slide,
        8.43,
        4.58,
        2.90,
        0.72,
        fill=BLUE,
        line=None,
    )
    add_text(
        slide,
        "VERIFY FINAL ANSWER",
        8.43,
        4.58,
        2.90,
        0.72,
        size=15,
        color=WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    add_text(
        slide,
        "PRESERVE   ·   GATE   ·   VERIFY",
        7.75,
        5.52,
        4.25,
        0.38,
        size=13,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
        margin=0,
    )

    # Original presentation footer.
    bottom_rule = add_connector(slide, 1.03, 7.02, 12.30, 7.02, width=1.0)
    bottom_rule.line.color.rgb = rgb(LINE)
    corner = slide.shapes.add_shape(
        MSO_SHAPE.PIE_WEDGE,
        Inches(11.99),
        Inches(6.68),
        Inches(0.58),
        Inches(0.58),
    )
    set_shape_fill(corner, CORNER)
    set_shape_line(corner, None)
    corner.rotation = 180

    prs.core_properties.title = "SEEK Motivation — Research Question and Solution"
    prs.core_properties.subject = "Thesis defense introduction slide"
    prs.core_properties.author = "SEEK"
    prs.save(PPTX_PATH)


# A coordinate-matched PNG preview for visual checking in environments without
# a PowerPoint renderer.
SCALE = 120
PX_W = round(SLIDE_W * SCALE)
PX_H = round(SLIDE_H * SCALE)

SERIF_REG = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
SERIF_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"
SANS_REG = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
SANS_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"


def px(value: float) -> int:
    return round(value * SCALE)


def hex_color(value: str) -> str:
    return f"#{value}"


def font(size: float, *, serif=False, bold=False):
    if serif:
        path = SERIF_BOLD if bold else SERIF_REG
    else:
        path = SANS_BOLD if bold else SANS_REG
    return ImageFont.truetype(path, round(size * SCALE / 72))


def rr(draw, x, y, w, h, fill, outline=None, width=1, radius=0.13):
    draw.rounded_rectangle(
        (px(x), px(y), px(x + w), px(y + h)),
        radius=px(radius),
        fill=hex_color(fill),
        outline=hex_color(outline) if outline else None,
        width=max(1, round(width * SCALE / 72)),
    )


def text_center(
    draw,
    text,
    x,
    y,
    w,
    h,
    *,
    size,
    color=TEXT,
    serif=False,
    bold=False,
    spacing=2,
):
    f = font(size, serif=serif, bold=bold)
    bbox = draw.multiline_textbbox((0, 0), text, font=f, spacing=spacing, align="center")
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    draw.multiline_text(
        (px(x + w / 2) - tw / 2, px(y + h / 2) - th / 2 - bbox[1]),
        text,
        font=f,
        fill=hex_color(color),
        spacing=spacing,
        align="center",
    )


def text_left(draw, text, x, y, *, size, color=TEXT, serif=False, bold=False):
    draw.text(
        (px(x), px(y)),
        text,
        font=font(size, serif=serif, bold=bold),
        fill=hex_color(color),
    )


def circle(draw, label, x, y, d, *, fill, outline=None, text_color=WHITE, size=15):
    draw.ellipse(
        (px(x), px(y), px(x + d), px(y + d)),
        fill=hex_color(fill),
        outline=hex_color(outline) if outline else None,
        width=2 if outline else 1,
    )
    text_center(
        draw,
        label,
        x,
        y,
        d,
        d,
        size=size,
        color=text_color,
        bold=True,
    )


def render_preview() -> None:
    image = Image.new("RGB", (PX_W, PX_H), hex_color(WHITE))
    draw = ImageDraw.Draw(image)

    text_left(draw, "Motivation", 0.40, 0.15, size=28, color=TEAL, serif=True, bold=True)
    draw.line((px(1.02), px(0.82), px(12.30), px(0.82)), fill=hex_color(LINE), width=2)
    draw.rectangle(
        (px(10.86), 0, px(13.33), px(0.62)),
        fill=hex_color(TAB_GRAY),
    )
    text_center(
        draw,
        "Introduction",
        10.96,
        0.10,
        2.24,
        0.40,
        size=18,
        color="000000",
        serif=True,
        bold=True,
    )

    left_x, panel_y, panel_w, panel_h = 0.68, 1.20, 5.55, 5.15
    right_x = 7.10
    for x in (left_x, right_x):
        rr(draw, x + 0.045, panel_y + 0.055, panel_w, panel_h, SHADOW)
        rr(draw, x, panel_y, panel_w, panel_h, WHITE, LIGHT_BORDER, radius=0.14)

    # Center arrow.
    ax, ay, aw, ah = px(6.38), px(3.45), px(0.45), px(0.43)
    draw.polygon(
        [
            (ax, ay + ah * 0.25),
            (ax + aw * 0.58, ay + ah * 0.25),
            (ax + aw * 0.58, ay),
            (ax + aw, ay + ah / 2),
            (ax + aw * 0.58, ay + ah),
            (ax + aw * 0.58, ay + ah * 0.75),
            (ax, ay + ah * 0.75),
        ],
        fill=hex_color(BLUE),
    )

    circle(draw, "1", 0.98, 1.47, 0.53, fill=NAVY)
    text_left(draw, "RESEARCH QUESTION", 1.67, 1.52, size=18, color=NAVY, serif=True, bold=True)
    circle(draw, "2", 7.40, 1.47, 0.53, fill=NAVY)
    text_left(draw, "SOLUTION — SEEK", 8.09, 1.52, size=18, color=NAVY, serif=True, bold=True)

    rr(draw, 2.30, 2.14, 2.28, 0.55, NAVY, radius=0.11)
    text_center(draw, "DECOMPOSITION", 2.30, 2.14, 2.28, 0.55, size=14, color=WHITE, bold=True)
    draw.line((px(3.44), px(2.69), px(3.44), px(2.90)), fill=hex_color(LINE), width=2)
    draw.line((px(2.00), px(2.90), px(4.88), px(2.90)), fill=hex_color(LINE), width=2)
    draw.line((px(2.00), px(2.90), px(2.00), px(3.06)), fill=hex_color(LINE), width=2)
    draw.line((px(4.88), px(2.90), px(4.88), px(3.06)), fill=hex_color(LINE), width=2)

    rr(draw, 1.12, 3.06, 2.35, 0.96, LIGHT_BLUE, "9FBCE0")
    circle(draw, "+", 1.28, 3.27, 0.46, fill=WHITE, outline=BLUE, text_color=BLUE, size=17)
    text_center(draw, "Focused\nretrieval", 1.89, 3.15, 1.42, 0.75, size=14, color=NAVY, bold=True)
    rr(draw, 3.62, 3.06, 2.14, 0.96, PALE_GRAY, "BCC6CE")
    circle(draw, "!", 3.77, 3.27, 0.46, fill=WHITE, outline=NAVY, text_color=NAVY, size=16)
    text_center(draw, "Context loss\nError propagation", 4.33, 3.14, 1.29, 0.76, size=10.5, bold=True)

    rr(draw, 1.10, 4.42, 4.72, 1.30, PALE_BLUE, BLUE, width=1.3)
    text_center(
        draw,
        "How can adaptive RAG benefit from\ndecomposition without trusting it blindly?",
        1.34,
        4.53,
        4.24,
        1.05,
        size=16.5,
        serif=True,
        bold=True,
    )

    rr(draw, 8.60, 2.13, 2.55, 0.54, NAVY, radius=0.11)
    text_center(draw, "ORIGINAL QUESTION", 8.60, 2.13, 2.55, 0.54, size=13.5, color=WHITE, bold=True)
    draw.line((px(9.88), px(2.67), px(9.88), px(2.89)), fill=hex_color(LINE), width=2)
    draw.line((px(8.50), px(2.89), px(11.25), px(2.89)), fill=hex_color(LINE), width=2)
    draw.line((px(8.50), px(2.89), px(8.50), px(3.06)), fill=hex_color(LINE), width=2)
    draw.line((px(11.25), px(2.89), px(11.25), px(3.06)), fill=hex_color(LINE), width=2)

    for x, title, subtitle, icon in (
        (7.43, "Direct path", "Preserve context", "↔"),
        (9.98, "Decomposed path", "Retrieve if uncertain", "?"),
    ):
        rr(draw, x, 3.06, 2.36, 1.04, PALE_BLUE, LIGHT_BORDER)
        circle(draw, icon, x + 0.13, 3.25, 0.44, fill=WHITE, outline=BLUE, text_color=BLUE, size=14)
        text_center(draw, title, x + 0.67, 3.16, 1.57, 0.34, size=13.5, color=NAVY, bold=True)
        text_center(draw, subtitle, x + 0.67, 3.52, 1.57, 0.34, size=10.8, color=MUTED)

    draw.line((px(8.61), px(4.10), px(8.61), px(4.35)), fill=hex_color(LINE), width=2)
    draw.line((px(11.16), px(4.10), px(11.16), px(4.35)), fill=hex_color(LINE), width=2)
    draw.line((px(8.61), px(4.35), px(11.16), px(4.35)), fill=hex_color(LINE), width=2)
    draw.line((px(9.88), px(4.35), px(9.88), px(4.58)), fill=hex_color(LINE), width=2)
    rr(draw, 8.43, 4.58, 2.90, 0.72, BLUE, radius=0.12)
    text_center(draw, "VERIFY FINAL ANSWER", 8.43, 4.58, 2.90, 0.72, size=15, color=WHITE, bold=True)
    text_center(draw, "PRESERVE   ·   GATE   ·   VERIFY", 7.75, 5.52, 4.25, 0.38, size=13, color=TEAL, bold=True)

    draw.line((px(1.03), px(7.02), px(12.30), px(7.02)), fill=hex_color(LINE), width=2)
    draw.pieslice(
        (px(11.99), px(6.68), px(12.57), px(7.26)),
        start=180,
        end=270,
        fill=hex_color(CORNER),
    )

    image.save(PNG_PATH, quality=95)


if __name__ == "__main__":
    build_pptx()
    render_preview()
    print(PPTX_PATH)
    print(PNG_PATH)
