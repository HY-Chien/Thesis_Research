#!/usr/bin/env python3
"""Visual comparison of probability-weighted and monotonic tail-leakage scores."""

from __future__ import annotations

import subprocess
from pathlib import Path

OUT_DIR = Path(__file__).resolve().parent
PNG_OUT = OUT_DIR / "tail_leakage_score_curves.png"
PDF_OUT = OUT_DIR / "tail_leakage_score_curves.pdf"
SVG_OUT = OUT_DIR / "tail_leakage_score_curves.svg"
PPTX_OUT = OUT_DIR / "tail_leakage_score_curves.pptx"

BACKGROUND = "#FAFAF9"
INK = "#18191A"
MUTED = "#626B70"
TITLE = "#762326"
GREEN = "#1E9C72"
GREEN_DARK = "#137455"
GREEN_LIGHT = "#DCEFE7"
ORANGE = "#C65B42"
ORANGE_LIGHT = "#F2D8CF"
PANEL = "#FFFFFF"
PANEL_LINE = "#CED4D8"


def add_panel_box(fig, bounds):
    from matplotlib.patches import FancyBboxPatch

    x, y, w, h = bounds
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        transform=fig.transFigure,
        boxstyle="round,pad=0.008,rounding_size=0.012",
        facecolor=PANEL,
        edgecolor=PANEL_LINE,
        linewidth=1.1,
        zorder=-1,
    )
    fig.patches.append(patch)


def build_figure():
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib.patches import FancyBboxPatch

    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "axes.titleweight": "bold",
            "axes.labelcolor": INK,
            "xtick.color": MUTED,
            "ytick.color": MUTED,
        }
    )

    fig = plt.figure(figsize=(13.333333, 7.5), dpi=144, facecolor=BACKGROUND)
    fig.text(
        0.055,
        0.925,
        r"Why Not Use $-m\log m$ as the Final Risk Score?",
        fontsize=23,
        fontweight="bold",
        color=TITLE,
        ha="left",
        va="center",
    )
    fig.text(
        0.057,
        0.865,
        "A leakage score should never decrease when more probability escapes top-5.",
        fontsize=13.5,
        color=MUTED,
        style="italic",
        ha="left",
    )
    fig.lines.append(
        plt.Line2D([0.055, 0.945], [0.835, 0.835], transform=fig.transFigure, color="#344E5C", lw=1.2)
    )

    left_box = (0.055, 0.285, 0.48, 0.505)
    right_box = (0.56, 0.285, 0.385, 0.505)
    add_panel_box(fig, left_box)
    add_panel_box(fig, right_box)
    ax1 = fig.add_axes([0.085, 0.35, 0.42, 0.355], facecolor=PANEL)
    ax2 = fig.add_axes([0.59, 0.35, 0.325, 0.355], facecolor=PANEL)

    leakage = np.linspace(0.0001, 0.98, 700)
    mass = 1.0 - leakage
    mlogm = -(mass * np.log(mass))
    tls = -np.log(mass)

    ax1.plot(leakage, mlogm, color=ORANGE, lw=3.8, label=r"$-m\log m$")
    ml_x, ml_y = 1.0 - 1.0 / np.e, 1.0 / np.e
    ax1.scatter([ml_x], [ml_y], s=60, color=ORANGE, zorder=5)
    ax1.annotate(
        r"peak: leakage $=1-1/e\approx0.632$",
        xy=(ml_x, ml_y),
        xytext=(0.43, 0.415),
        fontsize=11,
        color=ORANGE,
        arrowprops=dict(arrowstyle="->", color=ORANGE, lw=1.2),
    )
    ax1.axvspan(0.632, 0.98, color=ORANGE_LIGHT, alpha=0.42, zorder=0)
    ax1.text(
        0.79,
        0.64,
        "more leakage,\nbut score falls",
        transform=ax1.transAxes,
        ha="center",
        va="center",
        fontsize=10.5,
        fontweight="bold",
        color=TITLE,
    )
    ax1.set_xlim(0, 1.0)
    ax1.set_ylim(0, 0.45)
    ax1.set_xlabel(r"Tail leakage $\ell=1-m$   (more leakage $\rightarrow$)", fontsize=11.5)
    ax1.set_ylabel(r"$-m\log m$", fontsize=12)
    ax1.set_title("Probability weighting makes the score turn around", fontsize=15, color=INK, pad=12)
    ax1.grid(axis="y", linestyle="--", alpha=0.25)
    ax1.legend(loc="upper left", frameon=False, fontsize=12)
    ax1.spines[["top", "right"]].set_visible(False)

    ax2.plot(leakage, tls, color=GREEN, lw=3.8, label=r"$-\log m$")
    ax2.fill_between(leakage, 0, tls, color=GREEN_LIGHT, alpha=0.58)
    ax2.annotate(
        "keeps rising",
        xy=(0.86, float(-np.log(0.14))),
        xytext=(0.54, 2.72),
        fontsize=11.5,
        fontweight="bold",
        color=GREEN_DARK,
        arrowprops=dict(arrowstyle="->", color=GREEN_DARK, lw=1.4),
    )
    ax2.set_xlim(0, 1.0)
    ax2.set_ylim(0, 4.2)
    ax2.set_xlabel(r"Tail leakage $\ell=1-m$   (more leakage $\rightarrow$)", fontsize=11.5)
    ax2.set_ylabel(r"$-\log m$", fontsize=12)
    ax2.set_title("Tail-Leakage Surprisal stays monotonic", fontsize=15, color=INK, pad=12)
    ax2.grid(axis="y", linestyle="--", alpha=0.25)
    ax2.legend(loc="upper left", frameon=False, fontsize=12)
    ax2.spines[["top", "right"]].set_visible(False)

    formula_box = FancyBboxPatch(
        (0.12, 0.19),
        0.76,
        0.07,
        transform=fig.transFigure,
        boxstyle="round,pad=0.008,rounding_size=0.012",
        facecolor=ORANGE_LIGHT,
        edgecolor="none",
    )
    fig.patches.append(formula_box)
    fig.text(
        0.27,
        0.225,
        r"$-m\log m = m\times(-\log m)$",
        fontsize=17,
        fontweight="bold",
        color=TITLE,
        ha="center",
        va="center",
    )
    fig.text(
        0.67,
        0.225,
        "The factor m suppresses the score when retained mass is already very small.",
        fontsize=10.8,
        color=INK,
        ha="center",
        va="center",
    )

    takeaway = FancyBboxPatch(
        (0.08, 0.075),
        0.84,
        0.072,
        transform=fig.transFigure,
        boxstyle="round,pad=0.008,rounding_size=0.012",
        facecolor=GREEN_LIGHT,
        edgecolor="none",
    )
    fig.patches.append(takeaway)
    fig.text(
        0.50,
        0.111,
        r"Takeaway: decomposition finds the variable $m$; $-\log m$ gives one-way leakage severity.",
        fontsize=15.5,
        fontweight="bold",
        color=GREEN_DARK,
        ha="center",
        va="center",
    )
    fig.text(
        0.50,
        0.034,
        r"Design criterion: global monotonicity—not only fit within the observed $m$ range.",
        fontsize=9.5,
        color=MUTED,
        ha="center",
    )
    return fig


def save_pptx() -> None:
    """Create a fully editable PowerPoint slide using native shapes and charts."""
    import math
    import sys

    sys.path.insert(0, "/home/S113065528/.codex_tmp/python_pptx")
    from pptx import Presentation
    from pptx.chart.data import XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    def rgb(hex_color: str) -> RGBColor:
        value = hex_color.lstrip("#")
        return RGBColor.from_string(value)

    def add_text(
        slide,
        text: str,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        size: float,
        color: str = INK,
        bold: bool = False,
        italic: bool = False,
        align=PP_ALIGN.LEFT,
        rotation: float = 0,
        margin: float = 0.02,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.rotation = rotation
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(margin)
        frame.margin_right = Inches(margin)
        frame.margin_top = Inches(0.01)
        frame.margin_bottom = Inches(0.01)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = rgb(color)
        return box

    def add_round_rect(slide, x, y, w, h, fill, line=None, radius_adjust=0.12):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shape.line.fill.background()
        else:
            shape.line.color.rgb = rgb(line)
            shape.line.width = Pt(0.8)
        try:
            shape.adjustments[0] = radius_adjust
        except (IndexError, ValueError):
            pass
        return shape

    plot_x_factor = 0.13
    plot_y_factor = 0.045
    plot_w_factor = 0.83
    plot_h_factor = 0.78

    def add_xy_chart(slide, x, y, w, h, xs, ys, color, y_max, y_major):
        data = XyChartData()
        series = data.add_series("metric")
        for x_value, y_value in zip(xs, ys):
            series.add_data_point(float(x_value), float(y_value))
        chart_shape = slide.shapes.add_chart(
            XL_CHART_TYPE.XY_SCATTER_SMOOTH_NO_MARKERS,
            Inches(x),
            Inches(y),
            Inches(w),
            Inches(h),
            data,
        )
        chart = chart_shape.chart
        chart.has_legend = False
        chart.has_title = False
        chart.font.name = "Arial"
        chart.font.size = Pt(9)

        # Direct no-fill formatting allows editable shapes behind the chart
        # to provide the shaded regions without rasterizing the slide.
        for container in (chart._chartSpace, chart._chartSpace.chart.plotArea):
            sp_pr = container.find(qn("c:spPr"))
            if sp_pr is None:
                sp_pr = OxmlElement("c:spPr")
                container.append(sp_pr)
            for child in list(sp_pr):
                if child.tag in {
                    qn("a:noFill"),
                    qn("a:solidFill"),
                    qn("a:gradFill"),
                    qn("a:blipFill"),
                    qn("a:pattFill"),
                }:
                    sp_pr.remove(child)
            sp_pr.insert(0, OxmlElement("a:noFill"))

        plot_area = chart._chartSpace.chart.plotArea
        old_layout = plot_area.find(qn("c:layout"))
        if old_layout is not None:
            plot_area.remove(old_layout)
        layout = OxmlElement("c:layout")
        manual_layout = OxmlElement("c:manualLayout")
        for tag, value in (
            ("c:layoutTarget", "inner"),
            ("c:xMode", "factor"),
            ("c:yMode", "factor"),
            ("c:wMode", "factor"),
            ("c:hMode", "factor"),
            ("c:x", str(plot_x_factor)),
            ("c:y", str(plot_y_factor)),
            ("c:w", str(plot_w_factor)),
            ("c:h", str(plot_h_factor)),
        ):
            element = OxmlElement(tag)
            element.set("val", value)
            manual_layout.append(element)
        layout.append(manual_layout)
        plot_area.insert(0, layout)

        plotted = chart.series[0]
        plotted.format.line.color.rgb = rgb(color)
        plotted.format.line.width = Pt(2.6)

        x_axis = chart.category_axis
        x_axis.minimum_scale = 0.0
        x_axis.maximum_scale = 1.0
        x_axis.major_unit = 0.2
        x_axis.has_major_gridlines = False
        x_axis.tick_labels.font.name = "Arial"
        x_axis.tick_labels.font.size = Pt(8.5)
        x_axis.tick_labels.font.color.rgb = rgb(MUTED)
        x_axis.format.line.color.rgb = rgb("#AEB7BC")
        x_axis.format.line.width = Pt(0.8)

        y_axis = chart.value_axis
        y_axis.minimum_scale = 0.0
        y_axis.maximum_scale = y_max
        y_axis.major_unit = y_major
        y_axis.has_major_gridlines = True
        y_axis.major_gridlines.format.line.color.rgb = rgb("#E2E6E8")
        y_axis.major_gridlines.format.line.width = Pt(0.6)
        y_axis.tick_labels.font.name = "Arial"
        y_axis.tick_labels.font.size = Pt(8.5)
        y_axis.tick_labels.font.color.rgb = rgb(MUTED)
        y_axis.format.line.color.rgb = rgb("#AEB7BC")
        y_axis.format.line.width = Pt(0.8)
        return chart_shape

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BACKGROUND)

    add_text(
        slide,
        "Why Not Use −m log m as the Final Risk Score?",
        0.68,
        0.25,
        11.9,
        0.48,
        size=23,
        color=TITLE,
        bold=True,
    )
    add_text(
        slide,
        "A leakage score should never decrease when more probability escapes top-5.",
        0.70,
        0.75,
        11.6,
        0.30,
        size=13.5,
        color=MUTED,
        italic=True,
    )
    divider = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(0.68),
        Inches(1.14),
        Inches(12.60),
        Inches(1.14),
    )
    divider.line.color.rgb = rgb("#344E5C")
    divider.line.width = Pt(1.0)

    add_round_rect(slide, 0.52, 1.38, 6.38, 3.95, PANEL, PANEL_LINE)
    add_round_rect(slide, 7.16, 1.38, 5.65, 3.95, PANEL, PANEL_LINE)
    add_text(
        slide,
        "Probability weighting makes the score turn around",
        0.76,
        1.57,
        5.95,
        0.35,
        size=14.4,
        bold=True,
    )
    add_text(
        slide,
        "Tail-Leakage Surprisal stays monotonic",
        7.42,
        1.57,
        5.10,
        0.35,
        size=14.4,
        bold=True,
    )

    # Native line keys keep the formulas separately editable.
    key1 = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(0.83),
        Inches(2.01),
        Inches(1.16),
        Inches(2.01),
    )
    key1.line.color.rgb = rgb(ORANGE)
    key1.line.width = Pt(2.5)
    add_text(slide, "−m log m", 1.22, 1.87, 1.20, 0.28, size=11.3)

    key2 = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(7.47),
        Inches(2.01),
        Inches(7.80),
        Inches(2.01),
    )
    key2.line.color.rgb = rgb(GREEN)
    key2.line.width = Pt(2.5)
    add_text(slide, "−log m", 7.86, 1.87, 1.05, 0.28, size=11.3)

    leakage = [0.0001 + i * (0.98 - 0.0001) / 89 for i in range(90)]
    mass = [1.0 - value for value in leakage]
    mlogm = [-(m_value * math.log(m_value)) for m_value in mass]
    tls = [-math.log(m_value) for m_value in mass]

    # Editable native shading, positioned beneath the transparent charts.
    left_chart_box = (0.78, 2.10, 5.86, 2.38)
    right_chart_box = (7.42, 2.10, 5.08, 2.38)
    left_plot_x = left_chart_box[0] + left_chart_box[2] * plot_x_factor
    left_plot_y = left_chart_box[1] + left_chart_box[3] * plot_y_factor
    left_plot_w = left_chart_box[2] * plot_w_factor
    left_plot_h = left_chart_box[3] * plot_h_factor
    decline_x = left_plot_x + left_plot_w * (1.0 - 1.0 / math.e)
    decline_w = left_plot_w * (1.0 - (1.0 - 1.0 / math.e))
    decline_band = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(decline_x),
        Inches(left_plot_y),
        Inches(decline_w),
        Inches(left_plot_h),
    )
    decline_band.fill.solid()
    decline_band.fill.fore_color.rgb = rgb(ORANGE_LIGHT)
    decline_band.line.fill.background()
    decline_band.name = "Orange declining-region shading"

    right_plot_x = right_chart_box[0] + right_chart_box[2] * plot_x_factor
    right_plot_y = right_chart_box[1] + right_chart_box[3] * plot_y_factor
    right_plot_w = right_chart_box[2] * plot_w_factor
    right_plot_h = right_chart_box[3] * plot_h_factor
    scale_x = Inches(right_plot_w) / 1000
    scale_y = Inches(right_plot_h) / 1000
    area_builder = slide.shapes.build_freeform(0, 1000, scale=(scale_x, scale_y))
    area_vertices = [
        (
            value * 1000,
            max(0.0, min(1000.0, (1.0 - score / 4.2) * 1000)),
        )
        for value, score in zip(leakage, tls)
    ]
    area_vertices.append((0.98 * 1000, 1000))
    area_builder.add_line_segments(area_vertices, close=True)
    green_area = area_builder.convert_to_shape(
        Inches(right_plot_x),
        Inches(right_plot_y),
    )
    green_area.fill.solid()
    green_area.fill.fore_color.rgb = rgb(GREEN_LIGHT)
    green_area.line.fill.background()
    green_area.name = "Green area under -log m"

    left_chart_shape = add_xy_chart(
        slide, *left_chart_box, leakage, mlogm, ORANGE, 0.45, 0.05
    )
    right_chart_shape = add_xy_chart(
        slide, *right_chart_box, leakage, tls, GREEN, 4.2, 0.5
    )
    left_plot_group = slide.shapes.add_group_shape([decline_band, left_chart_shape])
    left_plot_group.name = "Left chart with aligned shading"
    right_plot_group = slide.shapes.add_group_shape([green_area, right_chart_shape])
    right_plot_group.name = "Right chart with aligned area fill"

    add_text(slide, "−m log m", 0.52, 2.63, 0.28, 1.18, size=11.5, rotation=270)
    add_text(slide, "−log m", 7.16, 2.69, 0.28, 1.03, size=11.5, rotation=270)
    add_text(
        slide,
        "Tail leakage  ℓ = 1 − m    (more leakage →)",
        1.24,
        4.70,
        4.95,
        0.30,
        size=10.8,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Tail leakage  ℓ = 1 − m    (more leakage →)",
        7.75,
        4.70,
        4.40,
        0.30,
        size=10.8,
        align=PP_ALIGN.CENTER,
    )

    # Editable peak/callout annotations.
    peak_dot = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(4.72),
        Inches(2.62),
        Inches(0.08),
        Inches(0.08),
    )
    peak_dot.fill.solid()
    peak_dot.fill.fore_color.rgb = rgb(ORANGE)
    peak_dot.line.fill.background()
    add_text(
        slide,
        "peak: leakage = 1 − 1/e ≈ 0.632",
        3.42,
        2.18,
        2.70,
        0.28,
        size=10.5,
        color=ORANGE,
        align=PP_ALIGN.CENTER,
    )
    peak_line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(4.76),
        Inches(2.47),
        Inches(4.76),
        Inches(2.64),
    )
    peak_line.line.color.rgb = rgb(ORANGE)
    peak_line.line.width = Pt(1.0)
    add_text(
        slide,
        "more leakage,\nbut score falls",
        4.75,
        3.12,
        1.40,
        0.52,
        size=10.5,
        color=TITLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    fall_line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(5.42),
        Inches(3.63),
        Inches(5.75),
        Inches(3.95),
    )
    fall_line.line.color.rgb = rgb(TITLE)
    fall_line.line.width = Pt(1.0)

    add_text(
        slide,
        "keeps rising",
        9.65,
        3.06,
        1.25,
        0.28,
        size=11,
        color=GREEN_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    rise_line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(10.63),
        Inches(3.40),
        Inches(11.15),
        Inches(3.81),
    )
    rise_line.line.color.rgb = rgb(GREEN_DARK)
    rise_line.line.width = Pt(1.0)

    add_round_rect(slide, 1.20, 5.52, 10.95, 0.56, ORANGE_LIGHT)
    add_text(
        slide,
        "−m log m = m × (−log m)",
        1.65,
        5.60,
        3.65,
        0.36,
        size=16.5,
        color=TITLE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "The factor m suppresses the value when retained mass is already very small.",
        5.35,
        5.60,
        6.30,
        0.36,
        size=10.5,
        align=PP_ALIGN.CENTER,
    )

    add_round_rect(slide, 0.80, 6.30, 11.70, 0.56, GREEN_LIGHT)
    add_text(
        slide,
        "Takeaway: decomposition finds the variable m; −log m gives one-way leakage severity.",
        1.05,
        6.39,
        11.20,
        0.34,
        size=15.0,
        color=GREEN_DARK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        "Design criterion: global monotonicity—not only fit within the observed m range.",
        2.20,
        7.00,
        8.95,
        0.24,
        size=9.2,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    prs.save(PPTX_OUT)


def main() -> None:
    import sys

    if "--pptx-only" in sys.argv:
        save_pptx()
        print(f"Wrote {PPTX_OUT}")
        return

    import matplotlib.pyplot as plt

    fig = build_figure()
    fig.savefig(PNG_OUT, dpi=144, facecolor=BACKGROUND)
    fig.savefig(PDF_OUT, facecolor=BACKGROUND)
    fig.savefig(SVG_OUT, facecolor=BACKGROUND)
    plt.close(fig)
    subprocess.run(
        [
            "/home/S113065528/miniconda3/bin/python3",
            str(Path(__file__).resolve()),
            "--pptx-only",
        ],
        check=True,
    )
    for path in (PNG_OUT, PDF_OUT, SVG_OUT, PPTX_OUT):
        print(f"Wrote {path}")


if __name__ == "__main__":
    main()
